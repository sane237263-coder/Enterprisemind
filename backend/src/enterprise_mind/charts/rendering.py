"""Deterministic PNG/PDF/XLSX/CSV rendering for frozen chart snapshots."""

from __future__ import annotations

import csv
import hashlib
import io
import math
import os
import textwrap
from collections.abc import Iterable
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any, cast

import fitz  # type: ignore[import-untyped]
from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.chart import BarChart, LineChart, PieChart, Reference  # type: ignore[import-untyped]
from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]
from PIL import Image, ImageDraw, ImageFont
from PIL.PngImagePlugin import PngInfo

from enterprise_mind.presentation import format_beijing_time

from .models import ChartFormat, ChartPoint, ChartSnapshot, ChartValidationError

WIDTH = 1600
HEIGHT = 900
PLOT_LEFT = 140
PLOT_RIGHT = 1510
PLOT_TOP = 110
PLOT_BOTTOM = 650
# Three wrapped category-label lines end at PLOT_BOTTOM + 16 + 2 * 18.
# Keep a deliberate gap before the provenance separator so labels cannot be
# mistaken for, or visually merge into, the source/metric footer.
FOOTER_TOP = 740
GRID_COLOR = (216, 224, 232)
INK = (33, 46, 60)
MUTED = (91, 108, 124)
ACCENT = (33, 104, 172)
ACCENT_ALT = (229, 126, 74)
ACCENT_THIRD = (68, 156, 116)
PIE_COLORS = [
    (33, 104, 172),
    (229, 126, 74),
    (68, 156, 116),
    (181, 93, 134),
    (164, 135, 58),
    (103, 94, 160),
    (64, 145, 150),
    (148, 91, 61),
]
X_LABEL_GAP = 8


@dataclass(frozen=True)
class RenderedArtifact:
    format: ChartFormat
    path: Path
    mime_type: str
    size: int
    sha256: str


def _font_candidates() -> Iterable[Path]:
    windows = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    yield from (
        windows / "msyh.ttc",
        windows / "msyh.ttf",
        windows / "simhei.ttf",
        windows / "simsun.ttc",
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    )


def _font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for candidate in _font_candidates():
        if candidate.is_file():
            try:
                return ImageFont.truetype(str(candidate), size)
            except OSError:
                continue
    return ImageFont.load_default()


def _draw_text(
    draw: ImageDraw.ImageDraw, xy: tuple[int, int], text: str, font: Any, fill: Any
) -> None:
    try:
        draw.text(xy, text, font=font, fill=fill)
    except (UnicodeEncodeError, ValueError):
        draw.text(xy, text.encode("ascii", "replace").decode("ascii"), font=font, fill=fill)


def _format_value(value: Decimal | None, precision: int | None) -> str:
    if value is None:
        return "NULL"
    if precision is None:
        return format(value, "f")
    quantum = Decimal(1).scaleb(-precision)
    return format(value.quantize(quantum, rounding=ROUND_HALF_UP), "f")


def _display_float(value: Decimal, precision: int | None) -> float:
    if precision is not None:
        quantum = Decimal(1).scaleb(-precision)
        value = value.quantize(quantum, rounding=ROUND_HALF_UP)
    result = float(value)
    if not math.isfinite(result):
        raise ChartValidationError(
            "CHART_VALUE_NONFINITE", "metric cannot be represented for drawing"
        )
    return result


def _title_lines(title: str, width: int = 44) -> list[str]:
    return textwrap.wrap(title, width=width, break_long_words=False, break_on_hyphens=False) or [
        title
    ]


def _footer_lines(snapshot: ChartSnapshot) -> list[str]:
    spec = snapshot.spec
    source_text = "; ".join(spec.source_refs)
    return [
        f"来源: {source_text}",
        f"统计口径: {spec.statistic}{('，单位: ' + spec.unit) if spec.unit else ''}",
        f"范围: {spec.scope_label}；查询时间: {format_beijing_time(spec.queried_at)}",
        f"数据引用: {spec.data_ref}；快照: {spec.content_hash}",
    ]


def _nice_step(value: float) -> float:
    if value <= 0 or not math.isfinite(value):
        return 1.0
    magnitude = 10.0 ** math.floor(math.log10(value))
    normalized = value / magnitude
    if normalized <= 1:
        multiple = 1.0
    elif normalized <= 2:
        multiple = 2.0
    elif normalized <= 5:
        multiple = 5.0
    else:
        multiple = 10.0
    return multiple * magnitude


def _axis_bounds(values: list[float]) -> tuple[float, float]:
    if not values:
        return 0.0, 1.0
    low = min(0.0, min(values))
    high = max(0.0, max(values))
    if low == high:
        padding = max(1.0, abs(high) * 0.1)
        low -= padding
        high += padding
    step = _nice_step((high - low) / 5.0)
    return math.floor(low / step) * step, math.ceil(high / step) * step


def _mapper(low: float, high: float) -> Any:
    span = high - low
    return lambda value: PLOT_BOTTOM - (value - low) / span * (PLOT_BOTTOM - PLOT_TOP)


def _draw_grid(draw: ImageDraw.ImageDraw, low: float, high: float, font: Any) -> Any:
    mapper = _mapper(low, high)
    step = _nice_step((high - low) / 5.0)
    tick_count = max(1, int(round((high - low) / step)))
    for index in range(tick_count + 1):
        value = high if index == tick_count else low + step * index
        y = int(mapper(value))
        draw.line((PLOT_LEFT, y, PLOT_RIGHT, y), fill=GRID_COLOR, width=2)
        _draw_text(draw, (20, y - 12), f"{value:g}", font, MUTED)
    zero_y = int(mapper(0.0))
    draw.line((PLOT_LEFT, zero_y, PLOT_RIGHT, zero_y), fill=INK, width=3)
    draw.line((PLOT_LEFT, PLOT_TOP, PLOT_LEFT, PLOT_BOTTOM), fill=INK, width=3)
    return mapper


def _category_label_lines(text: str) -> list[str]:
    return textwrap.wrap(text, width=12, break_long_words=False, break_on_hyphens=False)


def _text_bbox(
    draw: ImageDraw.ImageDraw, text: str, font: Any
) -> tuple[float, float, float, float]:
    try:
        return draw.textbbox((0, 0), text, font=font)
    except (UnicodeEncodeError, ValueError):
        fallback = text.encode("ascii", "replace").decode("ascii")
        return draw.textbbox((0, 0), fallback, font=font)


def _category_label_bounds(
    draw: ImageDraw.ImageDraw, text: str, x: int, font: Any
) -> tuple[float, float]:
    bounds = []
    for line in _category_label_lines(text)[:3]:
        bbox = _text_bbox(draw, line, font)
        left = int(x - (bbox[2] - bbox[0]) // 2) + bbox[0]
        bounds.append((left, left + bbox[2] - bbox[0]))
    if not bounds:
        return x, x
    return min(left for left, _ in bounds), max(right for _, right in bounds)


def _visible_x_label_indices(
    draw: ImageDraw.ImageDraw, rows: tuple[ChartPoint, ...], xs: list[int], font: Any
) -> list[int]:
    """Keep labels whose measured pixel bounds fit without overlap."""
    bounds = [
        _category_label_bounds(draw, row.dimension, x, font)
        for row, x in zip(rows, xs, strict=True)
    ]
    count = len(bounds)
    if count < 2 or all(
        bounds[index + 1][0] >= bounds[index][1] + X_LABEL_GAP for index in range(count - 1)
    ):
        return list(range(count))

    # Reserve room for the last label while walking forward. This keeps both
    # endpoints visible whenever their measured bounds allow it.
    selected = [0]
    last_left = bounds[-1][0]
    for index in range(1, count - 1):
        left, right = bounds[index]
        previous_right = bounds[selected[-1]][1]
        if left >= previous_right + X_LABEL_GAP and right + X_LABEL_GAP <= last_left:
            selected.append(index)
    if last_left >= bounds[selected[-1]][1] + X_LABEL_GAP:
        selected.append(count - 1)
    return selected


def _draw_labels(
    draw: ImageDraw.ImageDraw, rows: tuple[ChartPoint, ...], xs: list[int], font: Any
) -> None:
    visible_indices = _visible_x_label_indices(draw, rows, xs, font)
    for index in visible_indices:
        row, x = rows[index], xs[index]
        lines = _category_label_lines(row.dimension)
        for offset, line in enumerate(lines[:3]):
            bbox = _text_bbox(draw, line, font)
            _draw_text(
                draw,
                (int(x - (bbox[2] - bbox[0]) // 2), PLOT_BOTTOM + 16 + offset * 18),
                line,
                font,
                MUTED,
            )


def _draw_empty(
    draw: ImageDraw.ImageDraw, snapshot: ChartSnapshot, title_font: Any, body_font: Any
) -> None:
    _draw_text(draw, (PLOT_LEFT, 42), snapshot.spec.title, title_font, INK)
    _draw_text(draw, (PLOT_LEFT, 350), "无可绘制数据", _font(34), MUTED)
    _draw_text(draw, (PLOT_LEFT, 405), snapshot.spec.scope_label, body_font, MUTED)


def _draw_chart_image(snapshot: ChartSnapshot) -> Image.Image:
    spec = snapshot.spec
    image = Image.new("RGB", (WIDTH, HEIGHT), "white")
    draw = ImageDraw.Draw(image)
    title_font = _font(36)
    body_font = _font(20)
    small_font = _font(16)
    title_lines = _title_lines(spec.title)
    for index, line in enumerate(title_lines[:2]):
        _draw_text(draw, (PLOT_LEFT, 30 + index * 43), line, title_font, INK)
    if not snapshot.rows or not any(row.value is not None for row in snapshot.rows):
        _draw_empty(draw, snapshot, title_font, body_font)
    elif spec.chart_type in {"bar", "line"}:
        values = [
            _display_float(row.value, spec.display_precision)
            for row in snapshot.rows
            if row.value is not None
        ]
        low, high = _axis_bounds(values)
        mapper = _draw_grid(draw, low, high, small_font)
        count = len(snapshot.rows)
        step = (PLOT_RIGHT - PLOT_LEFT) / max(count, 1)
        xs = [int(PLOT_LEFT + step * (index + 0.5)) for index in range(count)]
        _draw_labels(draw, snapshot.rows, xs, small_font)
        if spec.chart_type == "bar":
            bar_width = max(8, int(step * 0.66))
            for row, x in zip(snapshot.rows, xs, strict=True):
                if row.value is None:
                    _draw_text(draw, (x - 18, PLOT_TOP + 8), "NULL", small_font, MUTED)
                    continue
                y = int(mapper(_display_float(row.value, spec.display_precision)))
                zero = int(mapper(0.0))
                draw.rectangle(
                    (x - bar_width // 2, min(y, zero), x + bar_width // 2, max(y, zero)),
                    fill=ACCENT,
                )
                label = _format_value(row.value, spec.display_precision)
                bbox = draw.textbbox((0, 0), label, font=small_font)
                _draw_text(
                    draw,
                    (int(x - (bbox[2] - bbox[0]) // 2), min(y, zero) - 24),
                    label,
                    small_font,
                    INK,
                )
        else:
            previous: tuple[int, int] | None = None
            for row, x in zip(snapshot.rows, xs, strict=True):
                if row.value is None:
                    previous = None
                    _draw_text(draw, (x - 18, PLOT_TOP + 8), "NULL", small_font, MUTED)
                    continue
                point = (x, int(mapper(_display_float(row.value, spec.display_precision))))
                if previous is not None:
                    draw.line((*previous, *point), fill=ACCENT, width=5)
                draw.ellipse((point[0] - 7, point[1] - 7, point[0] + 7, point[1] + 7), fill=ACCENT)
                previous = point
    else:
        values = [
            _display_float(row.value, spec.display_precision)
            for row in snapshot.rows
            if row.value is not None
        ]
        total = sum(values)
        center = (520, 390)
        radius = 250
        start = -90.0
        for index, row in enumerate(snapshot.rows):
            assert row.value is not None
            amount = _display_float(row.value, spec.display_precision)
            end = start + amount / total * 360.0
            draw.pieslice(
                (center[0] - radius, center[1] - radius, center[0] + radius, center[1] + radius),
                start,
                end,
                fill=PIE_COLORS[index % len(PIE_COLORS)],
                outline="white",
                width=3,
            )
            start = end
        for index, row in enumerate(snapshot.rows):
            y = 170 + index * 40
            color = PIE_COLORS[index % len(PIE_COLORS)]
            draw.rectangle((930, y, 956, y + 26), fill=color)
            percent = _display_float(cast(Decimal, row.value), spec.display_precision) / total * 100
            _draw_text(
                draw,
                (970, y),
                (
                    f"{row.dimension}: {percent:.1f}% "
                    f"({_format_value(row.value, spec.display_precision)})"
                ),
                body_font,
                INK,
            )

    footer_y = FOOTER_TOP
    draw.line((PLOT_LEFT, footer_y - 18, PLOT_RIGHT, footer_y - 18), fill=GRID_COLOR, width=2)
    for line in _footer_lines(snapshot):
        for wrapped in textwrap.wrap(
            line, width=145, break_long_words=False, break_on_hyphens=False
        ):
            _draw_text(draw, (PLOT_LEFT, footer_y), wrapped, small_font, MUTED)
            footer_y += 22
    return image


def _png_bytes(image: Image.Image, snapshot: ChartSnapshot) -> bytes:
    metadata = PngInfo()
    metadata.add_text("chart_type", snapshot.spec.chart_type)
    metadata.add_text("content_hash", snapshot.spec.content_hash or "")
    metadata.add_text("data_ref", snapshot.spec.data_ref)
    metadata.add_text("source_refs", ";".join(snapshot.spec.source_refs))
    metadata.add_text("statistic", snapshot.spec.statistic)
    metadata.add_text("queried_at", snapshot.spec.queried_at)
    metadata.add_text("scope", snapshot.spec.scope_label)
    stream = io.BytesIO()
    image.save(stream, format="PNG", optimize=False, compress_level=6, pnginfo=metadata)
    return stream.getvalue()


def _pdf_bytes(image: Image.Image, snapshot: ChartSnapshot) -> bytes:
    png = _png_bytes(image, snapshot)
    document = fitz.open()
    page = document.new_page(width=WIDTH, height=HEIGHT)
    page.insert_image(page.rect, stream=png)
    document.set_metadata(
        {
            "title": snapshot.spec.title,
            "author": "EnterpriseMind",
            "subject": (
                f"{snapshot.spec.chart_type}; {snapshot.spec.scope_label}; "
                f"{snapshot.spec.content_hash}"
            ),
            "keywords": ";".join(snapshot.spec.source_refs),
        }
    )
    output = cast(bytes, document.tobytes(garbage=4, deflate=True))
    document.close()
    return output


def _excel_metric(value: Decimal | None, precision: int | None) -> Decimal | None:
    if value is None:
        return None
    if precision is None:
        return value
    quantum = Decimal(1).scaleb(-precision)
    return value.quantize(quantum, rounding=ROUND_HALF_UP)


def _xlsx_bytes(snapshot: ChartSnapshot) -> bytes:
    spec = snapshot.spec
    workbook = Workbook()
    chart_sheet = workbook.active
    chart_sheet.title = "Chart"
    data_sheet = workbook.create_sheet("Data")
    source_sheet = workbook.create_sheet("Source")
    data_sheet.freeze_panes = "A2"
    headers = [
        spec.dimension_column,
        spec.metric_column,
        "metric_exact",
        "__source_refs",
        "__queried_at",
        "__statistic",
        "__scope",
        "__chart_type",
        "__content_hash",
    ]
    data_sheet.append(headers)
    for cell in data_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2168AC")
    for row in snapshot.rows:
        data_sheet.append(
            [
                row.dimension,
                _excel_metric(row.value, spec.display_precision),
                None if row.value is None else _format_value(row.value, spec.display_precision),
                ";".join(spec.source_refs),
                spec.queried_at,
                spec.statistic,
                spec.scope_label,
                spec.chart_type,
                spec.content_hash,
            ]
        )
    for row_number in range(2, len(snapshot.rows) + 2):
        data_sheet.cell(row_number, 2).number_format = "0.############################"
        data_sheet.cell(row_number, 3).number_format = "@"
    source_sheet.append(["field", "value"])
    source_sheet.append(["title", spec.title])
    source_sheet.append(["chart_type", spec.chart_type])
    source_sheet.append(["source_refs", ";".join(spec.source_refs)])
    source_sheet.append(["statistic", spec.statistic])
    source_sheet.append(["scope", spec.scope_label])
    source_sheet.append(["queried_at", spec.queried_at])
    source_sheet.append(["data_ref", spec.data_ref])
    source_sheet.append(["revision", spec.revision])
    source_sheet.append(["content_hash", spec.content_hash])
    source_sheet.append(["row_count", len(snapshot.rows)])
    source_sheet.append(["unit", spec.unit])
    for cell in source_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2168AC")
    if snapshot.rows and any(row.value is not None for row in snapshot.rows):
        if spec.chart_type == "bar":
            chart: Any = BarChart()
            chart.type = "col"
        elif spec.chart_type == "line":
            chart = LineChart()
            chart.display_blanks = "gap"
        else:
            chart = PieChart()
        data = Reference(data_sheet, min_col=2, min_row=1, max_row=len(snapshot.rows) + 1)
        categories = Reference(data_sheet, min_col=1, min_row=2, max_row=len(snapshot.rows) + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.title = spec.title
        chart.style = 10
        chart.height = 10
        chart.width = 20
        if spec.unit and spec.chart_type != "pie":
            chart.y_axis.title = spec.unit
        # Keep the native chart object on its own sheet.  The chart references
        # the typed values on Data and never embeds an external workbook URL.
        chart_sheet.add_chart(chart, "A1")
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _csv_text(value: str, field: str) -> str:
    if any(ord(char) < 32 and char not in "\t\n\r" for char in value):
        raise ChartValidationError(
            "CHART_CSV_CONTROL_CHARACTER", f"{field} contains a control character"
        )
    normalized = value.lstrip(" \t")
    if normalized.startswith(("=", "+", "-", "@")):
        raise ChartValidationError(
            "CHART_CSV_FORMULA_RISK", f"{field} begins with a formula-like character"
        )
    return value


def _csv_bytes(snapshot: ChartSnapshot) -> bytes:
    spec = snapshot.spec
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, lineterminator="\n")
    writer.writerow(
        [
            spec.dimension_column,
            spec.metric_column,
            "metric_exact",
            "__source_refs",
            "__queried_at",
            "__statistic",
            "__scope",
            "__chart_type",
            "__content_hash",
        ]
    )
    source_refs = ";".join(spec.source_refs)
    for row in snapshot.rows:
        writer.writerow(
            [
                _csv_text(row.dimension, "dimension"),
                "" if row.value is None else _format_value(row.value, spec.display_precision),
                "" if row.value is None else _format_value(row.value, spec.display_precision),
                _csv_text(source_refs, "source_refs"),
                _csv_text(spec.queried_at, "queried_at"),
                _csv_text(spec.statistic, "statistic"),
                _csv_text(spec.scope_label, "scope"),
                spec.chart_type,
                spec.content_hash,
            ]
        )
    return stream.getvalue().encode("utf-8-sig")


def _artifact(path: Path, chart_format: ChartFormat, mime_type: str) -> RenderedArtifact:
    data = path.read_bytes()
    return RenderedArtifact(
        format=chart_format,
        path=path,
        mime_type=mime_type,
        size=len(data),
        sha256=hashlib.sha256(data).hexdigest(),
    )


def render_chart(
    snapshot: ChartSnapshot,
    formats: Iterable[ChartFormat],
    output_dir: Path,
) -> dict[ChartFormat, RenderedArtifact]:
    """Render selected formats from one frozen snapshot without re-querying."""

    selected = tuple(dict.fromkeys(formats))
    unknown = set(selected) - set(("png", "pdf", "xlsx", "csv"))
    if unknown:
        raise ChartValidationError("CHART_FORMAT_INVALID", "unsupported chart format")
    target = output_dir.resolve()
    target.mkdir(parents=True, exist_ok=True)
    digest = snapshot.spec.content_hash
    if digest is None:
        raise ChartValidationError("CHART_HASH_REQUIRED", "snapshot must have a content hash")
    image: Image.Image | None = None
    artifacts: dict[ChartFormat, RenderedArtifact] = {}
    for chart_format in selected:
        path = target / f"chart-{digest[:16]}-{snapshot.spec.chart_type}.{chart_format}"
        if chart_format == "png":
            image = image or _draw_chart_image(snapshot)
            path.write_bytes(_png_bytes(image, snapshot))
            artifacts[chart_format] = _artifact(path, chart_format, "image/png")
        elif chart_format == "pdf":
            image = image or _draw_chart_image(snapshot)
            path.write_bytes(_pdf_bytes(image, snapshot))
            artifacts[chart_format] = _artifact(path, chart_format, "application/pdf")
        elif chart_format == "xlsx":
            path.write_bytes(_xlsx_bytes(snapshot))
            artifacts[chart_format] = _artifact(
                path,
                chart_format,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        else:
            path.write_bytes(_csv_bytes(snapshot))
            artifacts[chart_format] = _artifact(path, chart_format, "text/csv; charset=utf-8")
    return artifacts
