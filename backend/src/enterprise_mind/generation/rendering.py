"""Safe deterministic renderers for frozen Phase 9 content.

The renderer accepts only server-frozen text, structured tables, and already
authorized source summaries.  It never calls a model, opens a URL, executes
code, evaluates formulas, or treats a client path as an input.
"""

from __future__ import annotations

import asyncio
import csv
import hashlib
import importlib
import json
import math
import re
import unicodedata
import zipfile
from collections.abc import AsyncIterable, Awaitable, Callable, Iterable
from contextlib import suppress
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Protocol, cast

from enterprise_mind.api.errors import APIError, ErrorCode
from enterprise_mind.presentation import format_beijing_time

SUPPORTED_FORMATS = {"docx", "pdf", "md", "xlsx"}
MIMES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "md": "text/markdown; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv; charset=utf-8",
}
MAX_CELL_CHARS = 32767
MAX_COLUMNS = 16384
MAX_XLSX_ROWS = 1_048_576
_FORMULA_PREFIX = ("=", "+", "-", "@")
_CONTROL_RE = re.compile(r"[\x00-\x1f\x7f]")
_CONTROL_WITHOUT_NEWLINES_RE = re.compile(r"[\x00-\x09\x0b-\x1f\x7f]")


async def await_thread_completion[T](function: Callable[..., T], *args: Any, **kwargs: Any) -> T:
    """Run blocking work in a thread and defer cancellation until it exits.

    Cancelling ``asyncio.to_thread`` does not stop its underlying thread.  The
    caller must keep any files and streams used by that thread alive until it
    finishes, then propagate cancellation to the async caller.
    """

    operation = asyncio.create_task(asyncio.to_thread(function, *args, **kwargs))
    cancelled = False
    while not operation.done():
        try:
            await asyncio.shield(operation)
        except asyncio.CancelledError:
            cancelled = True
        except BaseException:
            # The thread can fail after the caller has already been cancelled.
            # Wait for it to finish, then preserve cancellation as the outcome.
            if not operation.done():
                raise
    if cancelled:
        with suppress(BaseException):
            operation.result()
        raise asyncio.CancelledError
    return operation.result()


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


@dataclass(frozen=True)
class RenderInput:
    title: str
    body_markdown: str
    tables: tuple[dict[str, Any], ...] = ()
    source_refs: tuple[dict[str, Any], ...] = ()


@dataclass(frozen=True)
class RenderedFile:
    path: Path
    file_format: str
    mime_type: str
    size_bytes: int
    content_sha256: str
    exported_rows: int | None = None
    executed_at: str | None = None
    # A chart render carries every artifact made from the same frozen stream.
    # Ordinary content and database exports leave this empty and continue to
    # use the single-file path.
    artifacts: tuple[RenderedFile, ...] = ()


def _safe_text(
    value: Any,
    *,
    field: str,
    max_chars: int = MAX_CELL_CHARS,
    preserve_newlines: bool = False,
) -> str:
    text_value: str = value if isinstance(value, str) else str(value)
    if preserve_newlines:
        text_value = text_value.replace("\r\n", "\n").replace("\r", "\n")
        text_value = _CONTROL_WITHOUT_NEWLINES_RE.sub("", text_value)
    else:
        text_value = _CONTROL_RE.sub("", text_value)
    if len(text_value) > max_chars:
        raise APIError(ErrorCode.FILE_TOO_LARGE, {"field": field})
    return text_value


def _safe_basename(value: str, extension: str) -> str:
    value = _safe_text(value, field="file_name", max_chars=500).strip()
    value = value.replace("\\", "/").rsplit("/", 1)[-1].strip(" .")
    value = value.removesuffix("." + extension).strip(" .")
    if not value:
        value = "enterprise-mind"
    return value + "." + extension


def _normal_tables(
    tables: Iterable[dict[str, Any]],
) -> list[tuple[str, list[str], list[list[Any]]]]:
    result: list[tuple[str, list[str], list[list[Any]]]] = []
    for raw in list(tables)[:32]:
        if not isinstance(raw, dict):
            raise APIError(ErrorCode.VALIDATION_ERROR)
        title = _safe_text(raw.get("title", "结果"), field="table.title", max_chars=200)
        columns = raw.get("columns")
        rows = raw.get("rows")
        if not isinstance(columns, list) or not isinstance(rows, list):
            raise APIError(ErrorCode.VALIDATION_ERROR)
        if not columns or len(columns) > MAX_COLUMNS:
            raise APIError(ErrorCode.VALIDATION_ERROR)
        names: list[str] = []
        used: dict[str, int] = {}
        for column in columns:
            name = _safe_text(column, field="table.column", max_chars=255).strip() or "列"
            used[name] = used.get(name, 0) + 1
            names.append(name if used[name] == 1 else f"{name} [{used[name]}]")
        normalized: list[list[Any]] = []
        if len(rows) + 1 > MAX_XLSX_ROWS:
            raise APIError(ErrorCode.FILE_TOO_LARGE, {"field": "table.rows"})
        for row in rows:
            if not isinstance(row, list) or len(row) != len(columns):
                raise APIError(ErrorCode.VALIDATION_ERROR)
            normalized.append(row)
        result.append((title, names, normalized))
    return result


def _source_lines(source_refs: Iterable[dict[str, Any]]) -> list[str]:
    lines = ["", "## 来源"]
    for index, source in enumerate(list(source_refs), 1):
        if not isinstance(source, dict):
            continue
        title = _safe_text(
            source.get("title") or source.get("source_title") or "来源",
            field="source.title",
            max_chars=500,
        )
        locator_value = source.get("locator") or source.get("database_name") or ""
        if source.get("referenced_relations"):
            locator_value = (
                f"{locator_value} · "
                + "、".join(str(item) for item in source["referenced_relations"])
            ).strip(" ·")
        if source.get("scope"):
            locator_value = f"{locator_value} · 范围 {source['scope']}".strip(" ·")
        source_time = source.get("executed_at") or source.get("queried_at")
        if source_time:
            label = "执行" if source.get("executed_at") else "查询时间"
            locator_value = (f"{locator_value} · {label} {format_beijing_time(source_time)}").strip(
                " ·"
            )
        locator = _safe_text(locator_value, field="source.locator", max_chars=1000)
        lines.append(f"{index}. {title}" + (f"（{locator}）" if locator else ""))
    return lines if len(lines) > 2 else []


def _display_text(value: Any, *, field: str) -> str:
    return "NULL" if value is None else _safe_text(value, field=field)


def _table_markdown(table: tuple[str, list[str], list[list[Any]]]) -> list[str]:
    title, columns, rows = table
    lines = [
        f"### {title}",
        "| " + " | ".join(columns) + " |",
        "| " + " | ".join("---" for _ in columns) + " |",
    ]
    for row in rows:
        lines.append(
            "| " + " | ".join(_display_text(value, field="table.cell") for value in row) + " |"
        )
    return lines


def _canonical_markdown(data: RenderInput) -> str:
    title = _safe_text(data.title, field="title", max_chars=300).strip() or "企业内容"
    body = _safe_text(
        data.body_markdown,
        field="body",
        max_chars=10_000_000,
        preserve_newlines=True,
    ).strip()
    if not body:
        raise APIError(ErrorCode.VALIDATION_ERROR)
    tables = _normal_tables(data.tables)
    sections = [f"# {title}", body]
    for table in tables:
        sections.extend(["", *_table_markdown(table)])
    sections.extend(_source_lines(data.source_refs))
    return "\n".join(sections).rstrip() + "\n"


def _excel_value(value: Any) -> Any:
    if value is None or isinstance(value, bool | int):
        return value
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, float):
        return value if math.isfinite(value) else _safe_text(value, field="cell")
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    text = _safe_text(value, field="cell")
    return "'" + text if text.startswith(_FORMULA_PREFIX) else text


def _render_markdown(path: Path, data: RenderInput) -> None:
    path.write_text(_canonical_markdown(data), encoding="utf-8", newline="\n")


def _render_docx(path: Path, data: RenderInput) -> None:
    from docx import Document
    from docx.enum.text import WD_BREAK

    document = Document()
    title = _safe_text(data.title, field="title", max_chars=300).strip() or "企业内容"
    document.add_heading(title, level=1)
    for raw_line in _safe_text(
        data.body_markdown,
        field="body",
        max_chars=10_000_000,
        preserve_newlines=True,
    ).splitlines():
        line = raw_line.strip()
        if not line:
            document.add_paragraph()
        elif line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("- ") or line.startswith("* "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("```"):
            paragraph = document.add_paragraph(style="No Spacing")
            paragraph.add_run(line).font.name = "Consolas"
            paragraph.add_run().add_break(WD_BREAK.LINE)
        else:
            document.add_paragraph(line)
    for table_data in _normal_tables(data.tables):
        title_text, columns, rows = table_data
        document.add_heading(title_text, level=2)
        table = document.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        for cell, value in zip(table.rows[0].cells, columns, strict=True):
            cell.text = value
        for row in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row, strict=True):
                cell.text = _display_text(value, field="table.cell")
    sources = _source_lines(data.source_refs)
    if sources:
        document.add_heading("来源", level=2)
        for line in sources[2:]:
            document.add_paragraph(line)
    document.save(str(path))


class _PdfFont(Protocol):
    @property
    def buffer(self) -> bytes: ...


class _PdfPage(Protocol):
    def insert_font(self, *, fontname: str, fontbuffer: bytes) -> int: ...

    def insert_textbox(
        self,
        rect: object,
        text: str,
        *,
        fontsize: float,
        lineheight: float,
        fontname: str,
        fontfile: str | None,
        color: tuple[float, float, float],
    ) -> float: ...


class _PdfDocument(Protocol):
    @property
    def page_count(self) -> int: ...

    def new_page(self) -> _PdfPage: ...

    def subset_fonts(self) -> None: ...

    def save(self, path: str, *, garbage: int, deflate: bool) -> None: ...

    def close(self) -> None: ...


class _FitzModule(Protocol):
    def open(self, path: str | None = None) -> _PdfDocument: ...

    def Rect(self, x0: float, y0: float, x1: float, y1: float) -> object: ...

    def Font(self, fontname: str) -> _PdfFont: ...


class _WorksheetWriter(Protocol):
    freeze_panes: str | None

    def append(self, row: Iterable[Any]) -> None: ...

    def close(self) -> None: ...


class _WorkbookWriter(Protocol):
    @property
    def worksheets(self) -> Iterable[_WorksheetWriter]: ...

    def create_sheet(self, title: str) -> _WorksheetWriter: ...

    def save(self, path: Path) -> None: ...

    def close(self) -> None: ...


class _OpenpyxlModule(Protocol):
    def Workbook(self, *, write_only: bool) -> _WorkbookWriter: ...


def _fitz_module() -> _FitzModule:
    return cast(_FitzModule, importlib.import_module("fitz"))


def _new_workbook() -> _WorkbookWriter:
    module = cast(_OpenpyxlModule, importlib.import_module("openpyxl"))
    return module.Workbook(write_only=True)


def _close_workbook(workbook: _WorkbookWriter) -> None:
    for worksheet in workbook.worksheets:
        with suppress(Exception):
            worksheet.close()
    with suppress(Exception):
        workbook.close()


def _pdf_lines(data: RenderInput) -> list[str]:
    title = _safe_text(data.title, field="title", max_chars=300).strip() or "企业内容"
    body = _safe_text(
        data.body_markdown,
        field="body",
        max_chars=10_000_000,
        preserve_newlines=True,
    ).strip()
    lines = [f"# {title}", *body.splitlines()]
    for table_title, columns, rows in _normal_tables(data.tables):
        lines.extend(("", f"表格：{table_title}", "列：" + " | ".join(columns)))
        for row in rows:
            lines.append("行：")
            for column, value in zip(columns, row, strict=True):
                lines.append(f"  {column}: {_display_text(value, field='table.cell')}")
    lines.extend(_source_lines(data.source_refs))
    return lines


def _wrap_pdf_line(line: str, max_width: float = 470.0) -> list[str]:
    if not line:
        return [""]
    chunks: list[str] = []
    current: list[str] = []
    current_width = 0.0
    for character in line:
        character_width = 10.0 if unicodedata.east_asian_width(character) in {"W", "F"} else 5.2
        if current and current_width + character_width > max_width:
            chunks.append("".join(current))
            current = []
            current_width = 0.0
        current.append(character)
        current_width += character_width
    if current:
        chunks.append("".join(current))
    return chunks


def _render_pdf(path: Path, data: RenderInput) -> None:
    fitz = _fitz_module()
    document = fitz.open()
    # The locked PyMuPDF wheel includes this embeddable CJK font. Using its
    # buffer avoids host-font differences and unsupported TTC subsetting in
    # Linux images, and never falls back to Latin-only Helvetica for Chinese.
    fontbuffer = fitz.Font("cjk").buffer
    fontname = "EMCJK"
    line_height = 14.0
    text_box_height = 20.0
    top = 54.0
    bottom = 788.0
    cursor_y = top
    page: _PdfPage | None = None
    try:
        for raw_line in _pdf_lines(data):
            for line in _wrap_pdf_line(raw_line):
                if page is None or cursor_y + line_height > bottom:
                    page = document.new_page()
                    page.insert_font(fontname=fontname, fontbuffer=fontbuffer)
                    cursor_y = top
                if line:
                    remaining = page.insert_textbox(
                        fitz.Rect(54, cursor_y, 541, cursor_y + text_box_height),
                        line,
                        fontsize=10,
                        lineheight=1.35,
                        fontname=fontname,
                        fontfile=None,
                        color=(0.12, 0.16, 0.22),
                    )
                    if remaining < 0:
                        raise APIError(ErrorCode.INTERNAL_ERROR)
                cursor_y += line_height
        if not document.page_count:
            document.new_page()
        # CJK fonts can make a one-page report nearly 20 MB when embedded whole.
        document.subset_fonts()
        document.save(str(path), garbage=4, deflate=True)
    finally:
        document.close()


def _render_xlsx(path: Path, data: RenderInput) -> None:
    workbook = _new_workbook()
    tables = _normal_tables(data.tables)
    body_lines = _safe_text(
        data.body_markdown,
        field="body",
        max_chars=10_000_000,
        preserve_newlines=True,
    ).splitlines()
    if tables:
        if body_lines:
            notes_sheet = workbook.create_sheet("说明")
            notes_sheet.append(["内容"])
            for line in body_lines:
                notes_sheet.append([_excel_value(line)])
        for title, columns, rows in tables:
            sheet = workbook.create_sheet(
                _safe_text(title, field="sheet", max_chars=31)[:31] or "结果"
            )
            sheet.freeze_panes = "A2"
            sheet.append([_excel_value(value) for value in columns])
            for row in rows:
                sheet.append([_excel_value(value) for value in row])
    else:
        sheet = workbook.create_sheet("正文")
        sheet.freeze_panes = "A2"
        sheet.append(["内容"])
        for line in body_lines:
            sheet.append([_excel_value(line)])
    sources = _source_lines(data.source_refs)
    if sources:
        source_sheet = workbook.create_sheet("来源")
        source_sheet.append(["来源"])
        for line in sources[2:]:
            source_sheet.append([_excel_value(line)])
    workbook.save(path)


def render_to_path(
    data: RenderInput,
    file_format: str,
    directory: Path,
    file_name: str | None = None,
    *,
    max_bytes: int = 100 * 1024 * 1024,
) -> RenderedFile:
    if file_format not in SUPPORTED_FORMATS:
        raise APIError(ErrorCode.UNSUPPORTED_FILE_TYPE)
    directory.mkdir(parents=True, exist_ok=True)
    name = _safe_basename(file_name or data.title, file_format)
    path = directory / name
    renderers = {
        "md": _render_markdown,
        "docx": _render_docx,
        "pdf": _render_pdf,
        "xlsx": _render_xlsx,
    }
    renderers[file_format](path, data)
    if not path.is_file() or path.stat().st_size == 0:
        raise APIError(ErrorCode.INTERNAL_ERROR)
    size = path.stat().st_size
    if size > max_bytes:
        raise APIError(ErrorCode.FILE_TOO_LARGE)
    if file_format == "xlsx" and not zipfile.is_zipfile(path):
        raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
    if file_format == "pdf":
        fitz = _fitz_module()
        document = fitz.open(str(path))
        try:
            if document.page_count < 1:
                raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
        finally:
            document.close()
    digest = _sha256_file(path)
    return RenderedFile(path, file_format, MIMES[file_format], size, digest)


async def _prepare_directory(directory: Path) -> None:
    await await_thread_completion(directory.mkdir, parents=True, exist_ok=True)


async def render_export_xlsx(
    frames: AsyncIterable[dict[str, Any]],
    directory: Path,
    file_name: str | None = None,
    *,
    source_refs: Iterable[dict[str, Any]] = (),
    max_rows: int = 1_000_000,
    max_bytes: int = 100 * 1024 * 1024,
    progress_callback: Callable[[int], Awaitable[None]] | None = None,
) -> RenderedFile:
    """Consume an authenticated export stream into one bounded XLSX file.

    The async iterator is consumed serially.  Only the current frame is held
    while openpyxl writes rows, so the complete database result never enters
    the model, a checkpoint, or an aggregate response object.
    """

    if not 1 <= max_rows <= MAX_XLSX_ROWS - 1:
        raise APIError(ErrorCode.VALIDATION_ERROR)
    await _prepare_directory(directory)
    path = directory / _safe_basename(file_name or "database-export", "xlsx")
    workbook: _WorkbookWriter | None = None
    try:
        workbook = _new_workbook()
        sheet: _WorksheetWriter | None = None
        column_count = 0
        expected_sequence = 1
        exported_rows = 0
        executed_at: str | None = None
        saw_header = False
        saw_end = False
        names: list[str] = []
        used_names: dict[str, int] = {}
        async for frame in frames:
            if not isinstance(frame, dict) or not isinstance(frame.get("type"), str):
                raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
            kind = frame["type"]
            if kind == "header":
                if saw_header or saw_end:
                    raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
                raw_columns = frame.get("columns")
                if not isinstance(raw_columns, list) or not raw_columns:
                    raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
                if len(raw_columns) > MAX_COLUMNS:
                    raise APIError(ErrorCode.FILE_TOO_LARGE)
                for column in raw_columns:
                    if not isinstance(column, dict):
                        raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
                    name = _safe_text(
                        column.get("name", "列"), field="export.column", max_chars=255
                    )
                    name = name.strip() or "列"
                    used_names[name] = used_names.get(name, 0) + 1
                    names.append(name if used_names[name] == 1 else f"{name} [{used_names[name]}]")
                sheet = workbook.create_sheet("导出结果")
                sheet.freeze_panes = "A2"
                sheet.append([_excel_value(name) for name in names])
                column_count = len(names)
                saw_header = True
            elif kind == "batch":
                if (
                    not saw_header
                    or sheet is None
                    or saw_end
                    or frame.get("sequence") != expected_sequence
                ):
                    raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
                rows = frame.get("rows")
                if not isinstance(rows, list):
                    raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
                for row in rows:
                    if not isinstance(row, list) or len(row) != column_count:
                        raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
                    exported_rows += 1
                    if exported_rows > max_rows:
                        raise APIError(ErrorCode.FILE_TOO_LARGE)
                    sheet.append([_excel_value(value) for value in row])
                if progress_callback is not None:
                    await progress_callback(exported_rows)
                expected_sequence += 1
            elif kind == "end":
                if (
                    not saw_header
                    or saw_end
                    or frame.get("complete") is not True
                    or frame.get("exported_rows") != exported_rows
                ):
                    raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
                if not isinstance(frame.get("executed_at"), str):
                    raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
                executed_at = _safe_text(
                    frame["executed_at"], field="export.executed_at", max_chars=100
                )
                saw_end = True
            else:
                raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
        if not saw_header or not saw_end or sheet is None:
            raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
        source_lines = _source_lines(
            ({**source, "executed_at": executed_at} for source in source_refs)
            if executed_at
            else source_refs
        )
        if source_lines:
            source_sheet = workbook.create_sheet("来源")
            source_sheet.append(["来源"])
            for line in source_lines[2:]:
                source_sheet.append([_excel_value(line)])
        workbook.save(path)
    except APIError:
        raise
    except Exception:
        raise APIError(ErrorCode.FILE_CONTENT_MISMATCH) from None
    finally:
        if workbook is not None:
            _close_workbook(workbook)
    if not path.is_file() or path.stat().st_size == 0:
        raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
    size = path.stat().st_size
    if size > max_bytes or not zipfile.is_zipfile(path):
        raise APIError(
            ErrorCode.FILE_TOO_LARGE if size > max_bytes else ErrorCode.FILE_CONTENT_MISMATCH
        )
    digest = _sha256_file(path)
    return RenderedFile(path, "xlsx", MIMES["xlsx"], size, digest, exported_rows, executed_at)


_CSV_NUMERIC_TYPES = {
    "int2",
    "int4",
    "int8",
    "smallint",
    "integer",
    "bigint",
    "numeric",
    "decimal",
    "float4",
    "float8",
    "real",
    "double precision",
    "money",
}


def _export_csv_cell(value: Any, data_type: str, *, field: str) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, dict | list):
        value = json.dumps(value, ensure_ascii=False, allow_nan=False, separators=(",", ":"))
    elif isinstance(value, Decimal):
        if not value.is_finite():
            raise APIError(ErrorCode.FILE_CONTENT_MISMATCH, {"field": field})
        value = format(value, "f")
    elif isinstance(value, float):
        if not math.isfinite(value):
            raise APIError(ErrorCode.FILE_CONTENT_MISMATCH, {"field": field})
        value = repr(value)
    elif isinstance(value, datetime | date | time):
        value = value.isoformat()
    else:
        value = str(value)
    if (
        len(value) > MAX_CELL_CHARS
        or any(ord(char) < 32 and char not in "\t\r\n" for char in value)
        or any(unicodedata.category(char) == "Cf" for char in value)
        or "\x7f" in value
    ):
        raise APIError(ErrorCode.FILE_CONTENT_MISMATCH, {"field": field})
    normalized = unicodedata.normalize("NFKC", value).lstrip()
    numeric = data_type.lower().split("(", 1)[0] in _CSV_NUMERIC_TYPES
    if numeric and normalized.lower() in {
        "nan",
        "+nan",
        "-nan",
        "inf",
        "+inf",
        "-inf",
        "infinity",
        "+infinity",
        "-infinity",
    }:
        raise APIError(ErrorCode.FILE_CONTENT_MISMATCH, {"field": field})
    if normalized.startswith(_FORMULA_PREFIX):
        if not numeric:
            raise APIError(ErrorCode.FILE_CONTENT_MISMATCH, {"field": field})
        try:
            if not Decimal(normalized).is_finite():
                raise ValueError
        except (ValueError, ArithmeticError):
            raise APIError(ErrorCode.FILE_CONTENT_MISMATCH, {"field": field}) from None
    return str(value)


async def render_export_csv(
    frames: AsyncIterable[dict[str, Any]],
    directory: Path,
    file_name: str | None = None,
    *,
    max_rows: int = 1_000_000,
    max_bytes: int = 100 * 1024 * 1024,
    progress_callback: Callable[[int], Awaitable[None]] | None = None,
) -> RenderedFile:
    """Write one bounded CSV from the authenticated database stream."""

    if not 1 <= max_rows <= 1_000_000 or max_bytes < 1:
        raise APIError(ErrorCode.VALIDATION_ERROR)
    await _prepare_directory(directory)
    path = directory / _safe_basename(file_name or "database-export", "csv")
    names: list[str] = []
    types: list[str] = []
    expected_sequence = 1
    exported_rows = 0
    executed_at: str | None = None
    saw_end = False
    try:
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream, lineterminator="\n")
            async for frame in frames:
                if not isinstance(frame, dict) or not isinstance(frame.get("type"), str):
                    raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
                kind = frame["type"]
                if kind == "header":
                    columns = frame.get("columns")
                    if names or saw_end or not isinstance(columns, list) or not columns:
                        raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
                    if len(columns) > MAX_COLUMNS:
                        raise APIError(ErrorCode.FILE_TOO_LARGE)
                    for column in columns:
                        if not isinstance(column, dict):
                            raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
                        name, data_type = column.get("name"), column.get("data_type", "")
                        if not isinstance(name, str) or not name or not isinstance(data_type, str):
                            raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
                        if name in names:
                            raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
                        names.append(_export_csv_cell(name, "text", field="export.column"))
                        types.append(data_type)
                    writer.writerow(names)
                elif kind == "batch":
                    if not names or saw_end or frame.get("sequence") != expected_sequence:
                        raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
                    rows = frame.get("rows")
                    if not isinstance(rows, list):
                        raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
                    for row in rows:
                        if not isinstance(row, list) or len(row) != len(names):
                            raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
                        exported_rows += 1
                        if exported_rows > max_rows:
                            raise APIError(ErrorCode.FILE_TOO_LARGE)
                        writer.writerow(
                            [
                                _export_csv_cell(value, data_type, field="export.cell")
                                for value, data_type in zip(row, types, strict=True)
                            ]
                        )
                    expected_sequence += 1
                    stream.flush()
                    if path.stat().st_size > max_bytes:
                        raise APIError(ErrorCode.FILE_TOO_LARGE)
                    if progress_callback is not None:
                        await progress_callback(exported_rows)
                elif kind == "end":
                    if (
                        not names
                        or saw_end
                        or frame.get("complete") is not True
                        or frame.get("exported_rows") != exported_rows
                        or not isinstance(frame.get("executed_at"), str)
                    ):
                        raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
                    executed_at = _safe_text(
                        frame["executed_at"], field="export.executed_at", max_chars=100
                    )
                    saw_end = True
                else:
                    raise APIError(ErrorCode.FILE_CONTENT_MISMATCH)
        if not saw_end or path.stat().st_size > max_bytes:
            code = ErrorCode.FILE_CONTENT_MISMATCH if not saw_end else ErrorCode.FILE_TOO_LARGE
            raise APIError(code)
        size = path.stat().st_size
        return RenderedFile(
            path, "csv", MIMES["csv"], size, _sha256_file(path), exported_rows, executed_at
        )
    except APIError:
        path.unlink(missing_ok=True)
        raise
    except Exception:
        path.unlink(missing_ok=True)
        raise APIError(ErrorCode.FILE_CONTENT_MISMATCH) from None


__all__ = [
    "MIMES",
    "RenderInput",
    "RenderedFile",
    "render_export_csv",
    "render_export_xlsx",
    "render_to_path",
]
