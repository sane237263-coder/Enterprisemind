"""Small generated files used by parser and service acceptance tests."""

import io

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches as SlideInches


def image_bytes(format="PNG"):
    stream = io.BytesIO()
    Image.new("RGB", (160, 100), "white").save(stream, format=format)
    return stream.getvalue()


def docx_bytes(*, pictures=False):
    document = Document()
    document.add_heading("设备规范", 1)
    document.add_paragraph("ABC-123 支持企业加密。<script>alert(1)</script>")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "访问控制"
    table.cell(1, 1).text = "部门"
    table.cell(2, 1).text = "公共"
    if pictures:
        document.add_picture(io.BytesIO(image_bytes()), width=Inches(1))
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def pptx_bytes(*, pictures=False):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text = slide.shapes.add_textbox(SlideInches(1), SlideInches(1), SlideInches(6), SlideInches(1))
    text.text = "ABC-123 设备架构"
    table = slide.shapes.add_table(
        3, 2, SlideInches(1), SlideInches(3), SlideInches(6), SlideInches(2)
    ).table
    table.cell(0, 0).text = "名称"
    table.cell(0, 1).text = "数量"
    table.cell(1, 0).merge(table.cell(2, 0))
    table.cell(1, 0).text = "设备"
    table.cell(1, 1).text = "3"
    table.cell(2, 1).text = "4"
    slide.notes_slide.notes_text_frame.text = "演讲备注：核对设备清单。"
    if pictures:
        slide.shapes.add_picture(io.BytesIO(image_bytes()), SlideInches(1), SlideInches(2))
    stream = io.BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def xlsx_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "设备清单"
    for row in [["型号", "数量"], ["ABC-123", 3], [None, 4], ["合计", "=SUM(B2:B3)"]]:
        sheet.append(row)
    sheet.merge_cells("A2:A3")
    sheet.row_dimensions[3].hidden = True
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def visual_pptx_bytes(*, smartart=False, chart=False, hidden=False, inherited=None):
    """A native vector flowchart or XY plot; SmartArt uses a detection-only graphic frame."""
    from pptx.chart.data import XyChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.oxml.xmlchemy import OxmlElement

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    if hidden:
        slide._element.set("show", "0")
        slide.shapes.add_textbox(
            SlideInches(1), SlideInches(1), SlideInches(6), SlideInches(1)
        ).text = "Hidden first slide"
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    if chart:
        values = XyChartData()
        series = values.add_series("ABC-123 observations")
        series.add_data_point(1, 2)
        series.add_data_point(2, 5)
        slide.shapes.add_chart(
            XL_CHART_TYPE.XY_SCATTER,
            SlideInches(1),
            SlideInches(1),
            SlideInches(6),
            SlideInches(4),
            values,
        )
    elif smartart:
        # This fixture verifies detection only. Real render acceptance uses upstream files.
        frame = OxmlElement("p:graphicFrame")
        nv = OxmlElement("p:nvGraphicFramePr")
        prop = OxmlElement("p:cNvPr")
        prop.set("id", "2")
        prop.set("name", "SmartArt")
        nv.append(prop)
        nv.append(OxmlElement("p:cNvGraphicFramePr"))
        nv.append(OxmlElement("p:nvPr"))
        frame.append(nv)
        xfrm = OxmlElement("p:xfrm")
        off = OxmlElement("a:off")
        off.set("x", "914400")
        off.set("y", "914400")
        ext = OxmlElement("a:ext")
        ext.set("cx", "5486400")
        ext.set("cy", "3657600")
        xfrm.append(off)
        xfrm.append(ext)
        frame.append(xfrm)
        graphic = OxmlElement("a:graphic")
        payload = OxmlElement("a:graphicData")
        payload.set("uri", "http://schemas.openxmlformats.org/drawingml/2006/diagram")
        graphic.append(payload)
        frame.append(graphic)
        slide.shapes._spTree.append(frame)
    else:
        first = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            SlideInches(1),
            SlideInches(2),
            SlideInches(2),
            SlideInches(1),
        )
        first.text = "ABC-123 Request"
        second = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, SlideInches(5), SlideInches(2), SlideInches(2), SlideInches(1)
        )
        second.text = "Approve"
        slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            SlideInches(3),
            SlideInches(2.5),
            SlideInches(5),
            SlideInches(2.5),
        )
    if inherited:
        assert not (smartart or chart or hidden)
        assert inherited in {"layout", "master"}
        target = slide.slide_layout if inherited == "layout" else slide.slide_layout.slide_master
        next_id = max(shape.shape_id for shape in target.shapes) + 1
        for number, shape in enumerate(list(slide.shapes)):
            shape._element.xpath(".//p:cNvPr")[0].set("id", str(next_id + number))
            target.shapes._spTree.insert_element_before(shape._element, "p:extLst")
    stream = io.BytesIO()
    presentation.save(stream)
    return stream.getvalue()
