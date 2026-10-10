import io
import json
from zipfile import ZIP_DEFLATED, ZipFile

import httpx
import pymupdf
import pytest
from defusedxml.common import DefusedXmlException
from fastapi.testclient import TestClient
from media_samples import docx_bytes, image_bytes, pptx_bytes, xlsx_bytes
from PIL import Image
from pydantic import ValidationError

from enterprise_mind.infrastructure.settings import UploadConfig, WorkerConfig
from enterprise_mind.ingestion.media_providers import HTTPOCR, HTTPVision
from enterprise_mind.ingestion.parsing import parse_file
from enterprise_mind.ingestion.uploads import UploadedFile, matches_signature
from enterprise_mind.office_converter import create_converter
from enterprise_mind.rag.chunking import AnchorInput, ElementInput, build_source
from enterprise_mind.rag.config import (
    ChunkConfig,
    OCRProviderConfig,
    ProvidersConfig,
    VisionProviderConfig,
)
from enterprise_mind.rag.providers import ProviderFailure


@pytest.mark.parametrize(
    "extension,builder", [("docx", docx_bytes), ("pptx", pptx_bytes), ("xlsx", xlsx_bytes)]
)
def test_office_text_tables_merge_anchors_and_escaped_preview(tmp_path, extension, builder):
    source = tmp_path / ("sample." + extension)
    source.write_bytes(builder())
    result = parse_file(source, extension, WorkerConfig(), tmp_path / "artifacts")
    assert "ABC-123" in "\n".join(s.text for s in result.sources)
    tables = [e for s in result.sources for e in s.elements if e.kind == "table"]
    assert tables and tables[0].payload["merged_ranges"] == ["A2:A3"]
    assert all(e.locator_type == extension for s in result.sources for e in s.elements)
    preview = (tmp_path / "artifacts/preview.html").read_text(encoding="utf-8")
    assert 'rowspan="2"' in preview and "<script>" not in preview
    if extension == "xlsx":
        assert "[公式未缓存]" in tables[0].text
        assert tables[0].payload["formulas"] == {"B4": "=SUM(B2:B3)"}
    if extension == "pptx":
        assert any(e.payload.get("notes") for e in result.sources[0].elements)


@pytest.mark.parametrize("extension,builder", [("docx", docx_bytes), ("pptx", pptx_bytes)])
def test_embedded_images_keep_original_object_location(tmp_path, extension, builder):
    path = tmp_path / ("sample." + extension)
    path.write_bytes(builder(pictures=True))
    parsed = parse_file(path, extension, WorkerConfig(), tmp_path / "artifacts")
    image = next(s for s in parsed.sources if s.kind == "image")
    assert image.metadata["needs_ocr"] and image.metadata["needs_vision"]
    assert image.metadata["origin_locator"]["object_id" if extension == "docx" else "shape_id"]
    assert any(a.artifact_type == "page_image" for a in parsed.artifacts)


@pytest.mark.parametrize(
    "extension,format", [("png", "PNG"), ("jpeg", "JPEG"), ("jpg", "JPEG"), ("webp", "WEBP")]
)
def test_image_types_signatures_and_pixel_limit(tmp_path, extension, format):
    body = image_bytes(format)
    assert matches_signature(extension, body[:16])
    path = tmp_path / ("sample." + extension)
    path.write_bytes(body)
    parsed = parse_file(path, extension, WorkerConfig())
    assert parsed.sources[0].width == 160
    with pytest.raises(ValueError):
        parse_file(path, extension, WorkerConfig(max_image_pixels=1024))
    file = UploadedFile("fake." + extension, "text/plain")
    file.validate_name(UploadConfig())
    assert str(file.error) == "FILE_CONTENT_MISMATCH"


def test_exif_orientation_matches_normalized_preview(tmp_path):
    path = tmp_path / "rotated.jpg"
    picture = Image.new("RGB", (120, 70))
    exif = Image.Exif()
    exif[274] = 6
    picture.save(path, exif=exif)
    result = parse_file(path, "jpg", WorkerConfig(), tmp_path / "artifacts")
    assert (result.sources[0].width, result.sources[0].height) == (70, 120)


def test_docx_nested_table_headers_and_footers_remain_searchable(tmp_path):
    from docx import Document

    document = Document()
    document.sections[0].header.paragraphs[0].text = "版本 ABC-123"
    document.sections[0].footer.paragraphs[0].text = "企业内部资料"
    outer = document.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "主表"
    inner = outer.cell(0, 0).add_table(rows=1, cols=2)
    inner.cell(0, 0).text, inner.cell(0, 1).text = "嵌套型号", "DEF-456"
    path = tmp_path / "nested.docx"
    document.save(path)
    result = parse_file(path, "docx", WorkerConfig(), tmp_path / "artifacts")
    text = result.sources[0].text
    assert all(value in text for value in ("ABC-123", "企业内部资料", "DEF-456"))
    assert any(e.locator.get("object_id") == "header-1" for e in result.sources[0].elements)


def test_pptx_native_chart_with_embedded_workbook(tmp_path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    chart = CategoryChartData()
    chart.categories = ["ABC-123", "DEF-456"]
    chart.add_series("数量", [3, 7])
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), chart
    )
    path = tmp_path / "chart.pptx"
    presentation.save(path)
    result = parse_file(path, "pptx", WorkerConfig(), tmp_path / "artifacts")
    table = next(e for e in result.sources[0].elements if e.kind == "table")
    assert "ABC-123" in table.text and "数量" in table.text and "7.0" in table.text
    assert table.locator["slide"] == 1 and table.locator["shape_id"]


@pytest.mark.parametrize(
    "name,content",
    [
        ("../escape.xml", b"<root/>"),
        ("word/vbaProject.bin", b"macro"),
        ("bad.xml", '<!DOCTYPE foo [<!ENTITY a "content">]><foo>&a;</foo>'.encode("utf-16")),
    ],
)
def test_packages_reject_traversal_macros_and_xml_entities(tmp_path, name, content):
    stream = io.BytesIO()
    with (
        ZipFile(io.BytesIO(docx_bytes())) as original,
        ZipFile(stream, "w", ZIP_DEFLATED) as modified,
    ):
        for item in original.infolist():
            modified.writestr(item.filename, original.read(item))
        modified.writestr(name, content)
    path = tmp_path / "bad.docx"
    path.write_bytes(stream.getvalue())
    with pytest.raises((ValueError, DefusedXmlException)):
        parse_file(path, "docx", WorkerConfig())
    with pytest.raises(ValueError):
        parse_file(path, "xlsx", WorkerConfig())


def test_native_pdf_table_and_rotated_ocr_image_coordinates(tmp_path):
    path = tmp_path / "table.pdf"
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=300, height=300)
        for x in (20, 140, 260):
            page.draw_line((x, 20), (x, 140))
        for y in (20, 60, 100, 140):
            page.draw_line((20, y), (260, y))
        for x, y, text in (
            (30, 45, "Item"),
            (150, 45, "Value"),
            (30, 85, "ABC-123"),
            (150, 85, "3"),
        ):
            page.insert_text((x, y), text)
        blank = pdf.new_page(width=300, height=200)
        blank.set_rotation(90)
        pdf.save(path)
    result = parse_file(path, "pdf", WorkerConfig(), tmp_path / "artifacts")
    assert any(e.kind == "table" for e in result.sources[0].elements)
    assert result.sources[1].metadata["needs_ocr"]
    artifact = next(a for a in result.artifacts if a.source_index == 1)
    assert artifact.metadata["height"] > artifact.metadata["width"]


def test_merged_rows_remain_together_when_chunking(tmp_path):
    path = tmp_path / "sample.xlsx"
    path.write_bytes(xlsx_bytes())
    table = parse_file(path, "xlsx", WorkerConfig()).sources[0].elements[0]
    element = ElementInput(
        1,
        None,
        0,
        "table",
        table.text,
        table.payload,
        AnchorInput(1, 1, 1, "xlsx", table.locator, table.source_text),
    )

    class Counter:
        def count(self, value):
            return len(value)

    drafts = build_source(
        [element],
        ChunkConfig(target_tokens=35, min_tokens=5, max_tokens=80, overlap_tokens=0),
        Counter(),
    )
    assert drafts
    assert any("ABC-123" in d.content and ",4\n" in d.content for d in drafts)
    assert all(d.metadata["merged_ranges"] == ["A2:A3"] for d in drafts)


def configure(tmp_path, monkeypatch, kind):
    secret = tmp_path / "key"
    secret.write_text("test-secret-not-for-production")
    monkeypatch.setenv(f"EM_{kind.upper()}_API_KEY_FILE", str(secret))
    values = {
        "provider_key": "test",
        "model_key": "test-v1",
        "model": "test",
        "base_url": "https://media.test",
        "max_retries": 1,
    }
    return OCRProviderConfig(**values) if kind == "ocr" else VisionProviderConfig(**values)


@pytest.mark.asyncio
async def test_ocr_contract_transient_retry_and_bounds(tmp_path, monkeypatch):
    config = configure(tmp_path, monkeypatch, "ocr")
    original, seen = httpx.AsyncClient, []

    def handler(request):
        seen.append(request)
        if len(seen) == 1:
            return httpx.Response(503)
        return httpx.Response(
            200,
            json={
                "blocks": [
                    {
                        "text": "ABC-123",
                        "box": {"x": 1.0, "y": 2.0, "width": 30.0, "height": 10.0},
                        "confidence": 0.98,
                    }
                ]
            },
        )

    monkeypatch.setattr(
        httpx,
        "AsyncClient",
        lambda **kwargs: original(**kwargs, transport=httpx.MockTransport(handler)),
    )
    provider = HTTPOCR(config)
    result = await provider.recognize(image_bytes(), 160, 100)
    assert result.blocks[0].text == "ABC-123" and len(seen) == 2
    assert json.loads(seen[-1].content)["mime_type"] == "image/png"
    with pytest.raises(ProviderFailure, match="OCR_RESPONSE_INVALID"):
        await provider.recognize(image_bytes(), 5, 5)


@pytest.mark.asyncio
async def test_vision_marks_semantics_and_rejects_truncation(tmp_path, monkeypatch):
    provider = HTTPVision(configure(tmp_path, monkeypatch, "vision"))

    async def response(path, body):
        assert body["messages"][0]["role"] == "system"
        return {"choices": [{"finish_reason": "length", "message": {"content": "{}"}}]}

    monkeypatch.setattr(provider, "request", response)
    with pytest.raises(ProviderFailure, match="VISION_RESPONSE_INVALID"):
        await provider.describe(image_bytes(), 160, 100)
    with pytest.raises(ValidationError):
        ProvidersConfig(allowed_origins=[], vision=provider.vision_config)


def test_converter_requires_auth_and_rejects_untrusted_format(tmp_path):
    secret = tmp_path / "token"
    secret.write_text("local-test-token")
    with TestClient(create_converter(secret)) as client:
        assert client.post("/convert/doc", content=b"fake").status_code == 401
        headers = {"Authorization": "Bearer local-test-token"}
        assert client.post("/convert/exe", content=b"MZ", headers=headers).status_code == 422
        assert client.post("/convert/doc", content=b"fake", headers=headers).status_code == 422
