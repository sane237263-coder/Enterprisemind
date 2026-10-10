from __future__ import annotations

import csv
from datetime import UTC, datetime
from decimal import Decimal
from types import SimpleNamespace

import fitz
import pytest
from openpyxl import load_workbook
from PIL import Image, ImageDraw

from enterprise_mind.charts import (
    ChartPoint,
    ChartSnapshot,
    ChartSpec,
    ChartValidationError,
    freeze_database_answer,
    render_chart,
)
from enterprise_mind.charts.rendering import (
    ACCENT,
    HEIGHT,
    PLOT_LEFT,
    PLOT_RIGHT,
    WIDTH,
    X_LABEL_GAP,
    _axis_bounds,
    _category_label_bounds,
    _font,
    _mapper,
    _visible_x_label_indices,
)

QUERIED_AT = "2026-09-22T10:00:00+08:00"


def make_spec(chart_type: str, **overrides) -> ChartSpec:
    values = {
        "chart_type": chart_type,
        "title": "企业部门指标图",
        "data_ref": "run:phase10/chart:001",
        "dimension_column": "department",
        "metric_column": "amount",
        "statistic": "SUM(amount) 按部门聚合",
        "scope": "all_aggregated",
        "scope_label": "全部匹配数据（按部门聚合）",
        "source_refs": ("db:fixture-source-001",),
        "queried_at": QUERIED_AT,
        "revision": "revision:001",
        "unit": "元",
        "sort_order": "input",
        "dimension_kind": "category",
        "aggregate_complete": True,
    }
    values.update(overrides)
    return ChartSpec(**values)


def make_snapshot(chart_type: str) -> ChartSnapshot:
    if chart_type == "line":
        spec = make_spec(
            chart_type,
            dimension_column="month",
            dimension_kind="time",
            sort_order="ascending",
            scope_label="全部匹配时间点（已按月份升序）",
            statistic="SUM(amount) 按月聚合",
        )
        rows = [
            ChartPoint("2026-03-01", Decimal("30.00")),
            ChartPoint("2026-01-01", Decimal("10.00")),
            ChartPoint("2026-02-01", Decimal("20.00")),
        ]
    else:
        spec = make_spec(chart_type)
        rows = [
            ChartPoint("华东", Decimal("10.00")),
            ChartPoint("华南", Decimal("20.00")),
            ChartPoint("华北", Decimal("30.00")),
        ]
    return ChartSnapshot.freeze(spec, rows)


@pytest.mark.parametrize("chart_type", ["bar", "line", "pie"])
@pytest.mark.parametrize("chart_format", ["png", "pdf", "xlsx", "csv"])
def test_three_chart_types_render_all_four_real_formats(tmp_path, chart_type, chart_format):
    snapshot = make_snapshot(chart_type)

    artifacts = render_chart(snapshot, [chart_format], tmp_path)

    artifact = artifacts[chart_format]
    assert artifact.path.is_file()
    assert artifact.size == artifact.path.stat().st_size > 0
    assert artifact.sha256
    if chart_format == "png":
        with Image.open(artifact.path) as image:
            assert image.size == (1600, 900)
            assert image.info["content_hash"] == snapshot.spec.content_hash
            assert image.info["source_refs"] == "db:fixture-source-001"
    elif chart_format == "pdf":
        with fitz.open(artifact.path) as document:
            assert document.page_count == 1
            assert document.metadata["author"] == "EnterpriseMind"
            assert "revision:001" not in document.metadata["subject"]
    elif chart_format == "xlsx":
        workbook = load_workbook(artifact.path, data_only=False)
        assert workbook.sheetnames == ["Chart", "Data", "Source"]
        chart_sheet = workbook["Chart"]
        data = workbook["Data"]
        assert len(chart_sheet._charts) == 1
        assert data["B2"].data_type == "n"
        assert data["C2"].value == "10.00"
        assert workbook["Source"]["B2"].value == "企业部门指标图"
        assert not workbook._external_links
    else:
        rows = list(csv.DictReader(artifact.path.read_text(encoding="utf-8-sig").splitlines()))
        assert len(rows) == 3
        assert rows[0]["__source_refs"] == "db:fixture-source-001"
        assert rows[0]["__scope"] == snapshot.spec.scope_label
        assert rows[0]["metric_exact"] == "10.00"


def test_empty_snapshot_has_explicit_empty_state_and_no_fake_xlsx_chart(tmp_path):
    snapshot = ChartSnapshot.freeze(make_spec("bar"), [])

    artifacts = render_chart(snapshot, ["png", "pdf", "xlsx", "csv"], tmp_path)

    with Image.open(artifacts["png"].path) as image:
        assert image.size == (1600, 900)
    with fitz.open(artifacts["pdf"].path) as document:
        assert document.page_count == 1
    workbook = load_workbook(artifacts["xlsx"].path)
    assert workbook.sheetnames == ["Chart", "Data", "Source"]
    assert len(workbook["Chart"]._charts) == 0
    assert list(csv.reader(artifacts["csv"].path.read_text(encoding="utf-8-sig").splitlines()))


def test_chart_footer_is_separated_from_three_line_category_labels(tmp_path):
    snapshot = ChartSnapshot.freeze(
        make_spec("bar"),
        [ChartPoint("very-long-category-label", Decimal("10"))],
    )
    artifact = render_chart(snapshot, ["png"], tmp_path)["png"]
    # The layout contract is deterministic: the footer separator is below the
    # last possible wrapped x-axis label line, with a visible safety gap.
    from enterprise_mind.charts.rendering import FOOTER_TOP, PLOT_BOTTOM

    assert FOOTER_TOP - (PLOT_BOTTOM + 16 + 2 * 18) >= 20
    with Image.open(artifact.path) as image:
        assert image.size == (1600, 900)


def test_null_is_preserved_for_bar_and_line_but_pie_rejects_it():
    rows = [ChartPoint("A", Decimal("1")), ChartPoint("B", None)]
    bar = ChartSnapshot.freeze(make_spec("bar"), rows)
    line = ChartSnapshot.freeze(
        make_spec(
            "line",
            dimension_column="day",
            dimension_kind="time",
            sort_order="ascending",
            scope_label="全部匹配时间点",
        ),
        [ChartPoint("2026-01-01", Decimal("1")), ChartPoint("2026-01-02", None)],
    )
    assert bar.rows[1].value is None and line.rows[1].value is None
    with pytest.raises(ChartValidationError, match="CHART_NULL_PIE_VALUE"):
        ChartSnapshot.freeze(make_spec("pie"), rows)


@pytest.mark.parametrize(
    ("values", "error"),
    [
        ([Decimal("-1"), Decimal("2")], "CHART_NEGATIVE_PIE_VALUE"),
        ([Decimal("0"), Decimal("0")], "CHART_PIE_TOTAL_ZERO"),
    ],
)
def test_pie_rejects_negative_and_all_zero_values(values, error):
    rows = [ChartPoint("A", values[0]), ChartPoint("B", values[1])]

    with pytest.raises(ChartValidationError, match=error):
        ChartSnapshot.freeze(make_spec("pie"), rows)


def test_duplicate_dimensions_and_budgets_are_rejected():
    duplicate = [ChartPoint("A", Decimal("1")), ChartPoint("A", Decimal("2"))]
    with pytest.raises(ChartValidationError, match="CHART_DUPLICATE_DIMENSION"):
        ChartSnapshot.freeze(make_spec("bar"), duplicate)

    too_many = [ChartPoint(str(index), Decimal("1")) for index in range(51)]
    with pytest.raises(ChartValidationError, match="CHART_BUDGET_EXCEEDED"):
        ChartSnapshot.freeze(make_spec("bar"), too_many)


def test_time_line_is_explicitly_sorted_and_missing_values_are_not_filled():
    snapshot = make_snapshot("line")

    assert [row.dimension for row in snapshot.rows] == [
        "2026-01-01",
        "2026-02-01",
        "2026-03-01",
    ]

    with pytest.raises(ChartValidationError, match="CHART_TIME_ORDER_REQUIRED"):
        ChartSnapshot.freeze(
            make_spec(
                "line",
                dimension_column="month",
                dimension_kind="time",
                scope_label="全部匹配时间点",
            ),
            [ChartPoint("2026-01-01", Decimal("1"))],
        )


def test_axis_bounds_use_zero_and_readable_ticks_for_positive_and_negative_values():
    assert _axis_bounds([10.0, 20.0, 30.0]) == (0.0, 30.0)
    assert _axis_bounds([-5.0, 10.0]) == (-5.0, 10.0)
    assert _axis_bounds([1e-10])[1] < 1.0


def test_x_axis_keeps_short_category_labels_at_normal_density():
    rows = tuple(ChartPoint(f"部门{index:02}", Decimal(index)) for index in range(1, 13))
    step = (PLOT_RIGHT - PLOT_LEFT) / len(rows)
    xs = [int(PLOT_LEFT + step * (index + 0.5)) for index in range(len(rows))]
    canvas = Image.new("RGB", (WIDTH, HEIGHT), "white")
    draw = ImageDraw.Draw(canvas)

    assert _visible_x_label_indices(draw, rows, xs, _font(16)) == list(range(12))


def test_dense_daily_labels_are_thinned_while_all_points_and_exports_remain(tmp_path):
    dimensions = [f"2026-01-{day:02}" for day in range(1, 29)]
    rows = [ChartPoint(day, Decimal(index + 1)) for index, day in enumerate(dimensions)]
    snapshot = ChartSnapshot.freeze(
        make_spec(
            "line",
            dimension_column="day",
            dimension_kind="time",
            sort_order="ascending",
            scope_label="全部匹配日期点（按日升序）",
            statistic="SUM(amount) 按日聚合",
        ),
        rows,
    )
    step = (PLOT_RIGHT - PLOT_LEFT) / len(snapshot.rows)
    xs = [int(PLOT_LEFT + step * (index + 0.5)) for index in range(len(snapshot.rows))]
    canvas = Image.new("RGB", (WIDTH, HEIGHT), "white")
    draw = ImageDraw.Draw(canvas)
    label_font = _font(16)
    visible = _visible_x_label_indices(draw, snapshot.rows, xs, label_font)

    assert visible[0] == 0 and visible[-1] == 27
    assert 8 <= len(visible) < 28
    label_bounds = [
        _category_label_bounds(draw, snapshot.rows[index].dimension, xs[index], label_font)
        for index in visible
    ]
    assert all(
        next_left >= current_right + X_LABEL_GAP
        for (_, current_right), (next_left, _) in zip(label_bounds, label_bounds[1:], strict=False)
    )
    artifacts = render_chart(snapshot, ["png", "pdf", "xlsx", "csv"], tmp_path)

    values = [float(row.value) for row in snapshot.rows if row.value is not None]
    low, high = _axis_bounds(values)
    mapper = _mapper(low, high)
    with Image.open(artifacts["png"].path) as image:
        assert image.size == (WIDTH, HEIGHT)
        for row, x in zip(snapshot.rows, xs, strict=True):
            assert row.value is not None
            point_y = int(mapper(float(row.value)))
            assert image.getpixel((x, point_y)) == ACCENT

    workbook = load_workbook(artifacts["xlsx"].path, data_only=False)
    data_rows = list(workbook["Data"].iter_rows(min_row=2, values_only=True))
    assert [(row[0], row[2]) for row in data_rows] == [
        (row.dimension, format(row.value, "f")) for row in snapshot.rows
    ]
    csv_rows = list(
        csv.DictReader(artifacts["csv"].path.read_text(encoding="utf-8-sig").splitlines())
    )
    assert [(row["day"], row["metric_exact"]) for row in csv_rows] == [
        (row.dimension, format(row.value, "f")) for row in snapshot.rows
    ]


def test_precision_requires_policy_and_xlsx_preserves_exact_text(tmp_path):
    spec = make_spec("bar", display_precision=2)
    value = Decimal("12345678901234567890.1234")
    snapshot = ChartSnapshot.freeze(spec, [ChartPoint("A", value)])
    artifact = render_chart(snapshot, ["xlsx"], tmp_path)["xlsx"]

    workbook = load_workbook(artifact.path, data_only=False)
    assert workbook["Data"]["C2"].value == "12345678901234567890.12"
    assert workbook["Data"]["B2"].data_type == "n"
    with pytest.raises(ChartValidationError, match="CHART_PRECISION_POLICY_REQUIRED"):
        ChartSnapshot.freeze(make_spec("bar"), [ChartPoint("A", value)])


def test_csv_rejects_formula_like_dimension_without_mutating_it(tmp_path):
    snapshot = ChartSnapshot.freeze(make_spec("bar"), [ChartPoint("=SUM(A1:A2)", Decimal("1"))])

    with pytest.raises(ChartValidationError, match="CHART_CSV_FORMULA_RISK"):
        render_chart(snapshot, ["csv"], tmp_path)


def _authorized_answer(*, query_kind="aggregate", has_more=False, relations=("public.sales",)):
    binding = SimpleNamespace(
        database_source_id=7,
        schema_snapshot_id=8,
        source_version=9,
        permission_fingerprint="permission-fingerprint-v1",
    )
    result = SimpleNamespace(
        rows=(
            {"department": "华东", "amount": "10.00"},
            {"department": "华南", "amount": "20.00"},
        ),
        has_more=has_more,
        referenced_relations=tuple(relations),
    )
    guarded = SimpleNamespace(referenced_relations=tuple(relations))
    return SimpleNamespace(
        prepared=SimpleNamespace(binding=binding, page_index=0),
        result=result,
        guarded_query=guarded,
        query_kind=query_kind,
        queried_at=datetime(2026, 9, 22, 10, 0, tzinfo=UTC),
    )


def _authority(answer, query_id, result_id=None):
    context = SimpleNamespace(
        query_id=query_id,
        owner_user_id=11,
        conversation_id=22,
        prepared=answer.prepared,
        guarded_query=answer.guarded_query,
        current_result_id=result_id,
        current_page_index=0,
    )
    page = None
    if result_id is not None:
        page = SimpleNamespace(
            query_id=query_id,
            result_id=result_id,
            owner_user_id=11,
            conversation_id=22,
            current_result_id=result_id,
            schema_snapshot_id=8,
            source_version=9,
            permission_fingerprint="permission-fingerprint-v1",
            page_index=0,
            has_more=answer.result.has_more,
            rows=answer.result.rows,
        )
    return {
        "query_id": query_id,
        "result_id": result_id,
        "owner_user_id": 11,
        "conversation_id": 22,
        "query_context": context,
        "page": page,
    }


def test_authorized_database_answer_derives_scope_and_sources_from_guarded_binding():
    answer = _authorized_answer()
    snapshot = freeze_database_answer(
        answer,
        **_authority(answer, "query-001"),
        chart_type="bar",
        title="授权聚合",
        scope="all_aggregated",
        dimension_column="department",
        metric_column="amount",
        statistic="SUM(amount) 按部门聚合",
    )

    assert snapshot.spec.aggregate_complete is True
    assert snapshot.spec.scope_label.startswith("全部匹配数据（服务端聚合")
    assert snapshot.spec.source_refs[0].startswith("database-source:7:snapshot:8:")
    assert snapshot.spec.source_refs[1] == "database-result:query-001:aggregate"


def test_authorized_database_answer_rejects_incomplete_all_scope():
    answer = _authorized_answer(query_kind="detail")
    with pytest.raises(ChartValidationError, match="CHART_AGGREGATION_INCOMPLETE"):
        freeze_database_answer(
            answer,
            **_authority(answer, "query-002"),
            chart_type="bar",
            title="不完整聚合",
            scope="all_aggregated",
            dimension_column="department",
            metric_column="amount",
            statistic="SUM(amount) 按部门聚合",
        )


def test_authorized_current_page_never_claims_complete_aggregation():
    answer = _authorized_answer(query_kind="detail", has_more=True)
    snapshot = freeze_database_answer(
        answer,
        **_authority(answer, "query-003", "result-003"),
        chart_type="bar",
        title="当前页",
        scope="current_page",
        dimension_column="department",
        metric_column="amount",
        statistic="当前页金额",
    )

    assert snapshot.spec.aggregate_complete is False
    assert snapshot.spec.scope_label.startswith("当前页")


def test_authorized_database_answer_rejects_guard_result_relation_mismatch():
    answer = _authorized_answer(relations=("public.other",))
    answer.guarded_query = SimpleNamespace(referenced_relations=("public.sales",))
    with pytest.raises(ChartValidationError, match="CHART_DATABASE_BINDING_INVALID"):
        freeze_database_answer(
            answer,
            **_authority(answer, "query-004"),
            chart_type="bar",
            title="关系不一致",
            scope="all_aggregated",
            dimension_column="department",
            metric_column="amount",
            statistic="SUM(amount)",
        )
