from zipfile import ZipFile

import pytest
from office_visual_samples import excel_visual_bytes, word_visual_bytes

from enterprise_mind.infrastructure.settings import WorkerConfig
from enterprise_mind.ingestion.office_visuals import docx_drawings, render_copy
from enterprise_mind.ingestion.parsing import parse_file
from enterprise_mind.ingestion.rich_parsing import check_package


@pytest.mark.parametrize("options", [{}, {"shape": True}, {"header": True}])
def test_word_native_visuals_keep_source_and_render_copy(tmp_path, options):
    path = tmp_path / "source.docx"
    original = word_visual_bytes(**options)
    path.write_bytes(original)
    result = parse_file(path, "docx", WorkerConfig())
    assert "Native paragraph" in result.sources[0].text
    visuals = [s for s in result.sources if s.metadata.get("needs_render")]
    assert len(visuals) == 1
    assert visuals[0].metadata["origin_locator"]["paragraph_ordinal"] >= 2
    target = tmp_path / "copy.docx"
    render_copy(path, "docx", visuals[0].metadata["render_unit"], target)
    check_package(target, "docx", WorkerConfig(), forbid_external_relationships=True)
    with ZipFile(target) as archive:
        assert b"Native paragraph" not in archive.read("word/document.xml")
        assert b"headerReference" not in archive.read("word/document.xml")
    assert path.read_bytes() == original


def test_excel_chart_sheet_and_cached_only_render_copy(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "source.xlsx"
    path.write_bytes(excel_visual_bytes())
    result = parse_file(path, "xlsx", WorkerConfig())
    visual = next(s for s in result.sources if s.metadata.get("needs_render"))
    assert visual.metadata["render_unit"] == 1
    assert visual.metadata["origin_locator"]["sheet"] == "Capacity"
    target = tmp_path / "copy.xlsx"
    render_copy(path, "xlsx", 1, target)
    book = load_workbook(target)
    try:
        assert book["Chart Data"].sheet_state == "hidden"
        assert book["Chart Data"]["B4"].value is None
        assert not book["Capacity"].print_area
        assert len(book["Capacity"]._charts) == 1
    finally:
        book.close()
    with pytest.raises(ValueError, match="visual sheet"):
        render_copy(path, "xlsx", 0, target)


def test_word_picture_inside_native_group_still_requires_rendering():
    from docx import Document
    from docx.oxml import parse_xml

    paragraph = Document().add_paragraph()
    namespaces = (
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
        'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup"'
    )
    for uri, content in (
        ("http://schemas.openxmlformats.org/drawingml/2006/picture", "<pic:pic/>"),
        (
            "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup",
            "<wpg:wgp><pic:pic/><wpg:grpSp/></wpg:wgp>",
        ),
    ):
        paragraph.add_run()._r.append(
            parse_xml(
                f'<w:drawing {namespaces}><a:graphic><a:graphicData uri="{uri}">'
                f"{content}</a:graphicData></a:graphic></w:drawing>"
            )
        )
    drawings = docx_drawings(paragraph)
    assert len(drawings) == 1 and b"wpg:wgp" in drawings[0].xml.encode()
