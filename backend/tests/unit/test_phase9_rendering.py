import asyncio
import csv
import threading
from pathlib import Path
from typing import Any

import fitz
import pytest
from docx import Document
from openpyxl import load_workbook

from enterprise_mind.api.errors import APIError, ErrorCode
from enterprise_mind.generation import rendering
from enterprise_mind.generation.rendering import (
    RenderInput,
    render_export_csv,
    render_export_xlsx,
    render_to_path,
)


def test_stream_database_csv_keeps_rows_and_execution_metadata(tmp_path: Path) -> None:
    async def frames():
        yield {
            "type": "header",
            "columns": [
                {"name": "订单号", "data_type": "text"},
                {"name": "金额", "data_type": "numeric"},
                {"name": "说明", "data_type": "text"},
            ],
        }
        yield {
            "type": "batch",
            "sequence": 1,
            "rows": [["001", "-12.50", "第一行\n第二行"], ["002", None, "中文"]],
        }
        yield {
            "type": "end",
            "exported_rows": 2,
            "complete": True,
            "executed_at": "2026-09-26T10:00:00Z",
        }

    output = asyncio.run(render_export_csv(frames(), tmp_path, "orders.csv"))

    assert output.file_format == "csv" and output.exported_rows == 2
    assert output.executed_at == "2026-09-26T10:00:00Z"
    assert output.mime_type.startswith("text/csv")
    with output.path.open(encoding="utf-8-sig", newline="") as stream:
        assert list(csv.reader(stream)) == [
            ["订单号", "金额", "说明"],
            ["001", "-12.50", "第一行\n第二行"],
            ["002", "", "中文"],
        ]


def test_stream_database_csv_rejects_formula_like_text(tmp_path: Path) -> None:
    async def frames():
        yield {"type": "header", "columns": [{"name": "备注", "data_type": "text"}]}
        yield {"type": "batch", "sequence": 1, "rows": [[" \t＝HYPERLINK(1)"]]}

    with pytest.raises(APIError) as error:
        asyncio.run(render_export_csv(frames(), tmp_path))
    assert error.value.code == ErrorCode.FILE_CONTENT_MISMATCH
    assert not list(tmp_path.glob("*.csv"))


def test_stream_database_csv_rejects_nonfinite_numeric_value(tmp_path: Path) -> None:
    async def frames():
        yield {"type": "header", "columns": [{"name": "金额", "data_type": "numeric"}]}
        yield {"type": "batch", "sequence": 1, "rows": [["NaN"]]}

    with pytest.raises(APIError) as error:
        asyncio.run(render_export_csv(frames(), tmp_path))
    assert error.value.code == ErrorCode.FILE_CONTENT_MISMATCH


def sample() -> RenderInput:
    return RenderInput(
        title="阶段九测试报告",
        body_markdown="## 结论\n中文正文与 2026-09-20 关键数字 235。\n\n- 仅展示资料内容",
        tables=(
            {
                "title": "结果",
                "columns": ["编码", "数值", "公式样文本"],
                "rows": [["001", 235, "=1+1"], ["002", None, "@text"]],
            },
        ),
        source_refs=({"title": "合成来源", "locator": "第 1 页"},),
    )


def test_all_phase9_render_formats_are_real_and_safe(tmp_path: Path) -> None:
    data = sample()
    outputs = {
        fmt: render_to_path(data, fmt, tmp_path, "phase9-report." + fmt)
        for fmt in ("md", "docx", "pdf", "xlsx")
    }

    markdown = outputs["md"].path.read_text(encoding="utf-8")
    assert "## 结论\n中文正文与 2026-09-20 关键数字 235。" in markdown
    assert "\n\n- 仅展示资料内容" in markdown
    assert "| 002 | NULL | @text |" in markdown

    document = Document(outputs["docx"].path)
    docx_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "中文正文与 2026-09-20 关键数字 235。" in docx_text
    assert "仅展示资料内容" in docx_text
    assert "1. 合成来源（第 1 页）" in docx_text
    assert "1.  1." not in docx_text
    docx_table_text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    assert "NULL" in docx_table_text

    pdf = fitz.open(outputs["pdf"].path)
    try:
        assert outputs["pdf"].path.stat().st_size < 2_000_000
        pdf_text = "\n".join(page.get_text() for page in pdf).replace("\xa0", " ")
        assert pdf.page_count == 1
        assert "中文正文与 2026-09-20 关键数字 235。" in pdf_text
        assert "表格：结果" in pdf_text
        assert "编码: 001" in pdf_text
        assert "数值: 235" in pdf_text
        assert "公式样文本: =1+1" in pdf_text
        assert "数值: NULL" in pdf_text
        assert "仅展示资料内容" in pdf_text
    finally:
        pdf.close()

    workbook = load_workbook(outputs["xlsx"].path, read_only=True, data_only=False)
    notes = list(workbook["说明"].values)
    assert notes[0] == ("内容",)
    assert notes[1] == ("## 结论",)
    assert notes[2] == ("中文正文与 2026-09-20 关键数字 235。",)
    assert notes[4] == ("'- 仅展示资料内容",)
    assert notes[4][0][1:] == "- 仅展示资料内容"
    sheet = workbook["结果"]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    assert rows[0] == ("编码", "数值", "公式样文本")
    assert rows[1] == ("001", 235, "'=1+1")
    assert rows[2][2] == "'@text"


def test_pdf_paginates_long_cjk_body_without_losing_content(tmp_path: Path) -> None:
    body = "\n".join(f"第 {index} 行：中文内容保留分页边界。" for index in range(90))
    output = render_to_path(
        RenderInput(title="长正文", body_markdown=body),
        "pdf",
        tmp_path,
        "long.pdf",
    )

    pdf = fitz.open(output.path)
    try:
        text = "\n".join(page.get_text() for page in pdf).replace("\xa0", " ")
        assert pdf.page_count > 1
        assert "第 0 行：中文内容保留分页边界。" in text
        assert "第 45 行：中文内容保留分页边界。" in text
        assert "第 89 行：中文内容保留分页边界。" in text
    finally:
        pdf.close()


def test_pdf_embeds_cjk_without_host_fonts(tmp_path: Path, monkeypatch: Any) -> None:
    real_is_file = Path.is_file

    def without_fonts(path: Path) -> bool:
        if path.suffix.lower() in {".ttc", ".ttf", ".otf"}:
            return False
        return real_is_file(path)

    monkeypatch.setattr(Path, "is_file", without_fonts)
    body = "中文结论：报销期限 17 天，上限 480 元。\n列：编号 | 金额 | 文本"
    output = render_to_path(
        RenderInput(
            title="跨平台中文报告",
            body_markdown=body,
            source_refs=({"title": "合成来源", "locator": "第 1 页"},),
        ),
        "pdf",
        tmp_path,
        "portable.pdf",
    )
    with fitz.open(output.path) as pdf:
        text = "\n".join(page.get_text() for page in pdf).replace("\xa0", " ")
        assert all(line in text for line in body.splitlines())
        assert "合成来源" in text
        fonts = pdf.get_page_fonts(0)
        assert fonts and all(pdf.extract_font(font[0])[3] for font in fonts)
    assert output.path.stat().st_size < 2_000_000


def test_render_filename_is_a_safe_basename(tmp_path: Path) -> None:
    output = render_to_path(sample(), "md", tmp_path, "../private\\report.md")
    assert output.path.parent == tmp_path
    assert output.path.name == "report.md"


def test_stream_export_writes_batches_and_preserves_execution_metadata(tmp_path: Path) -> None:
    async def frames():
        yield {
            "type": "header",
            "columns": [{"name": "订单号", "data_type": "text"}],
        }
        yield {"type": "batch", "sequence": 1, "rows": [["001"], ["002"]]}
        yield {
            "type": "end",
            "exported_rows": 2,
            "complete": True,
            "executed_at": "2026-09-20T10:00:00+08:00",
        }

    output = asyncio.run(
        render_export_xlsx(
            frames(),
            tmp_path,
            "orders.xlsx",
            source_refs=(
                {
                    "source_type": "database",
                    "source_title": "业务数据库",
                    "database_name": "enterprise",
                    "referenced_relations": ['"public"."orders"'],
                    "scope": "all_matching",
                },
            ),
        )
    )
    assert output.exported_rows == 2
    assert output.executed_at == "2026-09-20T10:00:00+08:00"
    workbook = load_workbook(output.path, read_only=True, data_only=False)
    assert list(workbook["导出结果"].values) == [("订单号",), ("001",), ("002",)]
    assert "执行 2026-09-20 10:00:00（北京时间）" in list(workbook["来源"].values)[1][0]
    workbook.close()


def test_stream_export_xlsx_keeps_formula_like_headers_as_text(tmp_path: Path) -> None:
    async def frames():
        yield {
            "type": "header",
            "columns": [
                {"name": "=1+1", "data_type": "text"},
                {"name": "+SUM(A1:A2)", "data_type": "text"},
                {"name": "-2+2", "data_type": "text"},
                {"name": "@SUM(A1:A2)", "data_type": "text"},
                {"name": "订单号", "data_type": "text"},
                {"name": "订单号", "data_type": "text"},
                {"name": "金额", "data_type": "numeric"},
            ],
        }
        yield {
            "type": "batch",
            "sequence": 1,
            "rows": [["=3+3", "+3+3", "-3+3", "@SUM(A1:A2)", "中文", "重复列", "12.50"]],
        }
        yield {
            "type": "end",
            "exported_rows": 1,
            "complete": True,
            "executed_at": "2026-09-29T00:00:00+00:00",
        }

    output = asyncio.run(render_export_xlsx(frames(), tmp_path, "headers.xlsx"))
    workbook = load_workbook(output.path, read_only=True, data_only=False)
    try:
        sheet = workbook["导出结果"]
        header = next(sheet.iter_rows(min_row=1, max_row=1))
        assert [cell.value for cell in header] == [
            "'=1+1",
            "'+SUM(A1:A2)",
            "'-2+2",
            "'@SUM(A1:A2)",
            "订单号",
            "订单号 [2]",
            "金额",
        ]
        assert all(cell.data_type != "f" for cell in header)
        row = next(sheet.iter_rows(min_row=2, max_row=2))
        assert [cell.value for cell in row] == [
            "'=3+3",
            "'+3+3",
            "'-3+3",
            "'@SUM(A1:A2)",
            "中文",
            "重复列",
            "12.50",
        ]
        assert all(cell.data_type != "f" for cell in row)
    finally:
        workbook.close()


def test_stream_export_rejects_missing_batch_sequence(tmp_path: Path) -> None:
    async def frames():
        yield {"type": "header", "columns": [{"name": "id"}]}
        yield {"type": "batch", "sequence": 2, "rows": [[1]]}

    try:
        asyncio.run(render_export_xlsx(frames(), tmp_path))
    except APIError as error:
        assert error.code == ErrorCode.FILE_CONTENT_MISMATCH
    else:
        raise AssertionError("invalid export sequence was accepted")


def test_stream_export_does_not_consume_frames_when_directory_prepare_fails(
    tmp_path: Path, monkeypatch: Any
) -> None:
    target = tmp_path / "unavailable"
    consumed = 0

    def fail_mkdir(path: Path, *args: Any, **kwargs: Any) -> None:
        raise OSError("directory preparation failed")

    monkeypatch.setattr(Path, "mkdir", fail_mkdir)

    async def frames():
        nonlocal consumed
        consumed += 1
        yield {"type": "header", "columns": [{"name": "id"}]}

    with pytest.raises(OSError, match="directory preparation failed"):
        asyncio.run(render_export_xlsx(frames(), target))

    assert consumed == 0
    assert not target.exists()


def test_stream_exports_keep_independent_task_directories_isolated(tmp_path: Path) -> None:
    left = tmp_path / "task-left"
    right = tmp_path / "task-right"
    headers_seen = 0
    headers_ready = asyncio.Event()

    def make_frames(value: str):
        async def frames():
            nonlocal headers_seen
            yield {"type": "header", "columns": [{"name": "id"}]}
            headers_seen += 1
            if headers_seen == 2:
                headers_ready.set()
            await headers_ready.wait()
            yield {"type": "batch", "sequence": 1, "rows": [[value]]}
            yield {"type": "end", "exported_rows": 1, "complete": True, "executed_at": "now"}

        return frames()

    async def run():
        return await asyncio.gather(
            render_export_xlsx(make_frames("left"), left, "export.xlsx"),
            render_export_xlsx(make_frames("right"), right, "export.xlsx"),
        )

    outputs = asyncio.run(run())
    assert {output.path.parent for output in outputs} == {left, right}

    left_workbook = load_workbook(outputs[0].path, read_only=True, data_only=False)
    right_workbook = load_workbook(outputs[1].path, read_only=True, data_only=False)
    try:
        assert list(left_workbook["导出结果"].values) == [("id",), ("left",)]
        assert list(right_workbook["导出结果"].values) == [("id",), ("right",)]
    finally:
        left_workbook.close()
        right_workbook.close()


def test_stream_export_works_with_an_existing_directory(tmp_path: Path) -> None:
    target = tmp_path / "existing"
    target.mkdir()

    async def frames():
        yield {"type": "header", "columns": [{"name": "id"}]}
        yield {"type": "batch", "sequence": 1, "rows": [["existing"]]}
        yield {"type": "end", "exported_rows": 1, "complete": True, "executed_at": "now"}

    output = asyncio.run(render_export_xlsx(frames(), target, "existing.xlsx"))

    assert output.path == target / "existing.xlsx"
    workbook = load_workbook(output.path, read_only=True, data_only=False)
    try:
        assert list(workbook["导出结果"].values) == [("id",), ("existing",)]
    finally:
        workbook.close()


def test_stream_export_closes_writer_on_validation_failure(
    tmp_path: Path, monkeypatch: Any
) -> None:
    real_workbook = rendering._new_workbook()
    closed = False

    class WorkbookSpy:
        @property
        def worksheets(self) -> Any:
            return real_workbook.worksheets

        def create_sheet(self, title: str) -> Any:
            return real_workbook.create_sheet(title)

        def save(self, path: Path) -> None:
            real_workbook.save(path)

        def close(self) -> None:
            nonlocal closed
            closed = True
            real_workbook.close()

    monkeypatch.setattr(rendering, "_new_workbook", lambda: WorkbookSpy())

    async def frames():
        yield {"type": "header", "columns": [{"name": "id"}]}
        yield {"type": "batch", "sequence": 2, "rows": [[1]]}

    try:
        asyncio.run(render_export_xlsx(frames(), tmp_path))
    except APIError as error:
        assert error.code == ErrorCode.FILE_CONTENT_MISMATCH
    else:
        raise AssertionError("invalid export sequence was accepted")
    assert closed


def test_stream_export_prepares_directory_in_worker_thread(
    tmp_path: Path, monkeypatch: Any
) -> None:
    target = tmp_path / "nested"
    caller_thread = threading.get_ident()
    mkdir_threads: list[int] = []
    original_mkdir = Path.mkdir

    def observed_mkdir(path: Path, *args: Any, **kwargs: Any) -> None:
        mkdir_threads.append(threading.get_ident())
        original_mkdir(path, *args, **kwargs)

    monkeypatch.setattr(Path, "mkdir", observed_mkdir)

    async def frames():
        yield {"type": "header", "columns": [{"name": "id"}]}
        yield {"type": "end", "exported_rows": 0, "complete": True, "executed_at": "now"}

    asyncio.run(render_export_xlsx(frames(), target))

    assert mkdir_threads
    assert mkdir_threads[0] != caller_thread


def test_stream_export_waits_for_directory_thread_after_cancellation(
    tmp_path: Path, monkeypatch: Any
) -> None:
    target = tmp_path / "cancelled"
    entered = threading.Event()
    release = threading.Event()
    original_mkdir = Path.mkdir

    def blocked_mkdir(path: Path, *args: Any, **kwargs: Any) -> None:
        entered.set()
        assert release.wait(timeout=5)
        original_mkdir(path, *args, **kwargs)

    monkeypatch.setattr(Path, "mkdir", blocked_mkdir)

    async def run() -> None:
        async def frames():
            yield {"type": "header", "columns": [{"name": "id"}]}
            yield {
                "type": "end",
                "exported_rows": 0,
                "complete": True,
                "executed_at": "now",
            }

        task = asyncio.create_task(render_export_xlsx(frames(), target))
        assert await asyncio.to_thread(entered.wait, 5)
        task.cancel()
        await asyncio.sleep(0)
        assert not task.done()
        task.cancel()
        await asyncio.sleep(0)
        assert not task.done()
        release.set()
        try:
            await task
        except asyncio.CancelledError:
            pass
        else:
            raise AssertionError("cancelled export completed successfully")

    asyncio.run(run())
    assert target.is_dir()
    assert not list(target.glob("*.xlsx"))
