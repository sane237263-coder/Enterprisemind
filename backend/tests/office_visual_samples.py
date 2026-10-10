"""Small native Office objects with text that is absent from ordinary document cells."""

import io

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def word_visual_bytes(*, shape=False, header=False):
    document = Document()
    document.add_paragraph("Native paragraph ABC-123.")
    owner = document.sections[0].header if header else document
    paragraph = owner.add_paragraph()
    if shape:
        drawing = parse_xml(
            '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml">'
            '<v:rect id="Visual1" style="width:180pt;height:72pt" fillcolor="#77aadd">'
            "<v:textbox><w:txbxContent><w:p><w:r><w:t>DRAWING-ONLY APPROVAL</w:t>"
            "</w:r></w:p></w:txbxContent></v:textbox></v:rect></w:pict>"
        )
    else:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        values = CategoryChartData()
        values.categories = ["Alpha", "Beta"]
        values.add_series("Units", [2, 5])
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, Inches(6), Inches(4), values
        ).chart
        chart.has_title = True
        chart.chart_title.text_frame.text = "CHART-ONLY CAPACITY"
        chart_part = Part(
            PackURI("/word/charts/chart1.xml"),
            chart.part.content_type,
            chart.part.blob,
            document.part.package,
        )
        workbook = next(iter(chart.part.rels.values())).target_part
        workbook_part = Part(
            PackURI("/word/embeddings/chart.xlsx"),
            workbook.content_type,
            workbook.blob,
            document.part.package,
        )
        assert chart_part.relate_to(workbook_part, RT.PACKAGE) == "rId1"
        relation = paragraph.part.relate_to(chart_part, RT.CHART)
        drawing = parse_xml(
            '<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<wp:inline><wp:extent cx="5486400" cy="3657600"/>'
            '<wp:docPr id="1" name="Native chart"/>'
            '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
            f'<c:chart r:id="{relation}"/></a:graphicData></a:graphic></wp:inline></w:drawing>'
        )
    paragraph.add_run()._r.append(drawing)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def excel_visual_bytes():
    book = Workbook()
    data = book.active
    data.title = "Chart Data"
    data.append(["Item", "Units"])
    data.append(["Alpha", 2])
    data.append(["Beta", 5])
    data["B4"] = '=WEBSERVICE("https://invalid.test/must-not-fetch")'
    sheet = book.create_sheet("Capacity")
    sheet["A1"] = "ABC-123 chart sheet"
    chart = BarChart()
    chart.title = "CHART-ONLY CAPACITY"
    chart.add_data(Reference(data, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(data, min_col=1, min_row=2, max_row=3))
    sheet.add_chart(chart, "D4")
    sheet.print_area = "A1:B2"  # The render copy must not clip the drawing to this old range.
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()
