"""Lifecycle contracts for database exports and chart bundles.

The PostgreSQL, Redis, and MinIO resources come from the isolated integration
fixtures. The real generation worker, XLSX writer, chart renderer, and storage
client are exercised. A real DatabaseMCPClient uses a controlled HTTP
transport so tests can provide exact MCP stream frames without starting or
claiming a real MCP process.
"""

from __future__ import annotations

import asyncio
import csv
import hashlib
import io
import json
import threading
import uuid
import zipfile
from collections.abc import AsyncIterator
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from datetime import UTC, datetime, timedelta
from typing import Any

import fitz
import httpx
import pytest
from openpyxl import load_workbook
from PIL import Image, ImageChops
from sqlalchemy import func, select
from sqlalchemy.orm import Session
from test_ingestion import ingestion_runtime  # noqa: F401

from enterprise_mind.agent.generation_models import GeneratedFile, GenerationTask
from enterprise_mind.agent.models import Conversation
from enterprise_mind.api.errors import ErrorCode
from enterprise_mind.database.catalog import catalog_to_wire
from enterprise_mind.database.export_models import DatabaseExportPlan
from enterprise_mind.database.mcp_client import DatabaseMCPClient
from enterprise_mind.database.models import (
    DatabaseSource,
    RelationPolicy,
    SchemaColumn,
    SchemaRelation,
    SchemaSnapshot,
)
from enterprise_mind.database.projection import authorized_schema
from enterprise_mind.identity.models import Department, User
from enterprise_mind.identity.service import CurrentIdentity
from enterprise_mind.infrastructure.settings import WorkerConfig
from enterprise_mind.ingestion.models import OutboxEvent, StorageObject

pytestmark = pytest.mark.integration


_COLUMNS = [
    {"name": "category", "type": "text"},
    {"name": "amount", "type": "numeric"},
]
_ROWS = [["north", 12], ["south", 7], ["east", 4]]
_EXECUTED_AT = "2026-09-25T00:00:00+00:00"


class _ControlledBody(httpx.AsyncByteStream):
    def __init__(
        self,
        payload: bytes,
        *,
        entered: threading.Event | None = None,
        release: threading.Event | None = None,
    ) -> None:
        self.payload = payload
        self.entered = entered
        self.release = release

    async def __aiter__(self) -> AsyncIterator[bytes]:
        if self.entered is not None:
            self.entered.set()
        if self.release is not None and not await asyncio.to_thread(self.release.wait, 30):
            raise TimeoutError("the controlled MCP stream was not released")
        yield self.payload

    async def aclose(self) -> None:
        return None


class _ControlledMCPTransport(httpx.AsyncBaseTransport):
    """Protocol-shaped MCP HTTP responses with explicit plan ownership."""

    def __init__(
        self,
        handle: str,
        relation: str,
        *,
        stream_error: str | None = None,
        initially_registered: bool = True,
        release_http_failures: int = 0,
        release_connect_failure: bool = False,
        stream_barriers: tuple[tuple[threading.Event, threading.Event], ...] = (),
    ) -> None:
        self.handle = handle
        self.relation = relation
        self.stream_error = stream_error
        self.active_handles = {handle} if initially_registered else set()
        self.release_http_failures = release_http_failures
        self.release_connect_failure = release_connect_failure
        self.stream_barriers = stream_barriers
        self.stream_requests: list[dict[str, Any]] = []
        self.release_requests: list[dict[str, Any]] = []

    async def handle_async_request(self, request: httpx.Request) -> httpx.Response:
        payload = json.loads(await request.aread())
        if request.url.path.endswith("/export-plan/stream"):
            self.stream_requests.append(payload)
            assert payload["handle"] == self.handle
            if self.handle not in self.active_handles:
                frames = [{"type": "error", "error_code": ErrorCode.QUERY_CONTEXT_EXPIRED.value}]
            else:
                frames = [
                    {
                        "type": "header",
                        "columns": _COLUMNS,
                        "referenced_relations": [self.relation],
                    },
                    {"type": "batch", "sequence": 1, "rows": _ROWS},
                ]
                if self.stream_error is not None:
                    frames.append({"type": "error", "error_code": self.stream_error})
                else:
                    frames.append(
                        {
                            "type": "end",
                            "exported_rows": len(_ROWS),
                            "complete": True,
                            "executed_at": _EXECUTED_AT,
                        }
                    )
            body = b"".join(
                json.dumps(frame, separators=(",", ":")).encode("utf-8") + b"\n" for frame in frames
            )
            call_index = len(self.stream_requests) - 1
            barrier = (
                self.stream_barriers[call_index] if call_index < len(self.stream_barriers) else None
            )
            return httpx.Response(
                200,
                headers={"Content-Type": "application/x-ndjson"},
                stream=_ControlledBody(
                    body,
                    entered=barrier[0] if barrier else None,
                    release=barrier[1] if barrier else None,
                ),
                request=request,
            )

        if request.url.path.endswith("/export-plan/release"):
            self.release_requests.append(payload)
            assert payload["handle"] == self.handle
            if self.release_connect_failure:
                raise httpx.ConnectError("controlled release transport failure", request=request)
            if self.release_http_failures:
                self.release_http_failures -= 1
                return httpx.Response(503, content=b"unavailable", request=request)
            self.active_handles.discard(self.handle)
            return httpx.Response(200, json={"ok": True}, request=request)

        return httpx.Response(404, content=b"not found", request=request)

    async def aclose(self) -> None:
        return None


def _create_export_task(runtime: Any, kind: str, *, file_format: str = "xlsx") -> dict[str, Any]:
    assert kind in {"database_export", "chart_render"}
    assert file_format in {"xlsx", "csv"}
    task_key = f"{kind}-{uuid.uuid4().hex}"
    handle = uuid.uuid4().hex

    async def create() -> dict[str, Any]:
        async with runtime.worker.sessions() as session, session.begin():
            row = (
                await session.execute(
                    select(User, Department)
                    .outerjoin(Department, User.department_id == Department.id)
                    .where(User.username == "初始管理员")
                )
            ).one_or_none()
            assert row is not None
            user, department = row
            identity = CurrentIdentity(user, department, user.credential_version)
            conversation = Conversation(owner_user_id=user.id, title="Export lifecycle " + task_key)
            session.add(conversation)

            source = DatabaseSource(
                name="Lifecycle source " + task_key,
                host="controlled-mcp.invalid",
                port=5432,
                database_name="lifecycle-business",
                ssl_mode="disable",
                secret_ref="secret://controlled-test",
                status="active",
            )
            session.add(source)
            await session.flush()
            snapshot = SchemaSnapshot(
                database_source_id=source.id,
                snapshot_no=1,
                status="active",
                captured_at=datetime.now(UTC),
            )
            session.add(snapshot)
            await session.flush()
            relation = SchemaRelation(
                snapshot_id=snapshot.id,
                schema_name="biz",
                relation_name="orders",
                relation_type="table",
                comment="Lifecycle test relation",
                relation_oid=701001,
                relation_signature="lifecycle-orders-signature",
            )
            session.add(relation)
            await session.flush()
            session.add_all(
                [
                    SchemaColumn(
                        relation_id=relation.id,
                        ordinal=1,
                        column_name="category",
                        data_type="text",
                    ),
                    SchemaColumn(
                        relation_id=relation.id,
                        ordinal=2,
                        column_name="amount",
                        data_type="numeric",
                    ),
                    RelationPolicy(
                        database_source_id=source.id,
                        schema_name="biz",
                        relation_name="orders",
                        enabled=True,
                    ),
                ]
            )
            source.active_snapshot_id = snapshot.id
            await session.flush()

            binding = await authorized_schema(session, identity, snapshot_id=int(snapshot.id))
            qualified_relation = binding.catalog.relations[0].quoted_qualified_name
            scope = "all_matching"
            chart_scope = "all_aggregated"
            source_refs = [
                {
                    "source_type": "database",
                    "source_title": binding.source_name,
                    "database_name": binding.database_name,
                    "database_source_id": str(binding.database_source_id),
                    "schema_snapshot_id": str(binding.schema_snapshot_id),
                    "source_version": binding.source_version,
                    "permission_fingerprint": binding.permission_fingerprint,
                    "referenced_relations": [qualified_relation],
                    "scope": chart_scope if kind == "chart_render" else scope,
                    "query_kind": "aggregate" if kind == "chart_render" else "detail",
                    "conditions_summary": "three controlled categories",
                    "queried_at": None,
                }
            ]
            plan = DatabaseExportPlan(
                owner_user_id=user.id,
                conversation_id=conversation.id,
                query_id=uuid.uuid4().hex,
                result_id=None,
                runtime_handle=handle,
                scope=scope,
                format="xlsx" if kind == "chart_render" else file_format,
                database_source_id=binding.database_source_id,
                schema_snapshot_id=binding.schema_snapshot_id,
                source_version=binding.source_version,
                permission_fingerprint=binding.permission_fingerprint,
                catalog=catalog_to_wire(binding.catalog),
                referenced_relations=[qualified_relation],
                conditions_summary="three controlled categories",
                expires_at=datetime.now(UTC) + timedelta(minutes=15),
            )
            session.add(plan)
            await session.flush()

            input_ref: dict[str, Any] = {
                "runtime_handle": handle,
                "file_name": "lifecycle-"
                + task_key
                + ("." + file_format if kind == "database_export" else ""),
            }
            task_format = file_format
            if kind == "chart_render":
                task_format = "png"
                input_ref.update(
                    {
                        "formats": ["png", "pdf", "xlsx", "csv"],
                        "query_id": plan.query_id,
                        "result_id": None,
                        "chart": {
                            "chart_type": "bar",
                            "title": "Amount by category",
                            "scope": chart_scope,
                            "dimension_column": "category",
                            "metric_column": "amount",
                            "statistic": "sum",
                            "dimension_kind": "category",
                            "sort_order": "input",
                            "top_n": None,
                        },
                    }
                )
            task = GenerationTask(
                owner_user_id=user.id,
                conversation_id=conversation.id,
                kind=kind,
                format=task_format,
                scope=scope,
                export_plan_id=plan.id,
                idempotency_key=task_key,
                request_fingerprint=hashlib.sha256(task_key.encode()).hexdigest(),
                input_ref=input_ref,
                source_refs=source_refs,
                status="queued",
            )
            session.add(task)
            await session.flush()
            event = OutboxEvent(
                aggregate_type="generation_task",
                aggregate_id=task.id,
                event_type="generation.export.requested",
                payload={"task_id": str(task.id)},
                deduplication_key="generation:" + str(task.id),
            )
            session.add(event)
            await session.flush()
            policy = await session.scalar(
                select(RelationPolicy).where(
                    RelationPolicy.database_source_id == source.id,
                    RelationPolicy.relation_name == "orders",
                )
            )
            assert policy is not None
            return {
                "task_id": int(task.id),
                "event_id": int(event.id),
                "plan_id": int(plan.id),
                "handle": handle,
                "owner_id": int(user.id),
                "source_id": int(source.id),
                "policy_id": int(policy.id),
                "relation": qualified_relation,
            }

    return runtime.run(create)


def test_database_csv_export_worker_and_private_download(runtime, database, ingestion_settings):
    task_data = _create_export_task(runtime, "database_export", file_format="csv")
    worker = runtime.worker.generation
    with _controlled_database(runtime, task_data) as (_worker, transport):
        assert runtime.run(worker.handle_event, task_data["event_id"]) is True
        state = _task_state(database, task_data["task_id"], task_data["plan_id"])
        assert state["status"] == "succeeded"
        assert state["exported_rows"] == len(_ROWS)
        assert state["plan_exists"] is False
        assert len(state["artifacts"]) == 1
        artifact = state["artifacts"][0]
        assert artifact["file_format"] == "csv"
        assert artifact["file_name"].endswith(".csv")
        assert len(transport.stream_requests) == 1
        assert len(transport.release_requests) == 1

    response = runtime.client.get(
        f"/api/v1/generated-files/{artifact['file_id']}/download", headers=runtime.admin
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("text/csv")
    assert response.content.startswith(b"\xef\xbb\xbf")
    assert list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")))) == [
        ["category", "amount"],
        ["north", "12"],
        ["south", "7"],
        ["east", "4"],
    ]


@contextmanager
def _controlled_database(
    runtime: Any,
    task_data: dict[str, Any],
    *,
    stream_error: str | None = None,
    initially_registered: bool = True,
    release_http_failures: int = 0,
    release_connect_failure: bool = False,
    stream_barriers: tuple[tuple[threading.Event, threading.Event], ...] = (),
) -> Any:
    worker = runtime.worker.generation
    previous_database = worker.database
    previous_plans = worker.export_plans
    transport = _ControlledMCPTransport(
        task_data["handle"],
        task_data["relation"],
        stream_error=stream_error,
        initially_registered=initially_registered,
        release_http_failures=release_http_failures,
        release_connect_failure=release_connect_failure,
        stream_barriers=stream_barriers,
    )
    client = DatabaseMCPClient(
        "http://controlled-mcp.invalid/mcp",
        "controlled-token",
        transport=transport,
    )
    worker.database = client
    # Force the durable DatabaseExportPlan path and the real MCP client stream.
    worker.export_plans = None
    try:
        yield worker, transport
    finally:
        runtime.run(client.close)
        worker.database = previous_database
        worker.export_plans = previous_plans


def _track_attempt_keys(storage: Any, monkeypatch: pytest.MonkeyPatch) -> list[str]:
    original = storage.new_key
    keys: list[str] = []

    def tracked(prefix: str = "original", **kwargs: Any) -> str:
        key = original(prefix, **kwargs)
        keys.append(key)
        return key

    monkeypatch.setattr(storage, "new_key", tracked)
    return keys


def _pause_after_uploads(
    storage: Any,
    monkeypatch: pytest.MonkeyPatch,
    expected_uploads: int,
) -> tuple[threading.Event, threading.Event, list[int]]:
    original = storage.upload
    entered = threading.Event()
    release = threading.Event()
    completed = [0]

    def upload_then_wait(key: str, stream: Any, mime: str) -> None:
        original(key, stream, mime)
        completed[0] += 1
        if completed[0] == expected_uploads:
            entered.set()
            if not release.wait(timeout=30):
                raise TimeoutError("the test did not release the real MinIO upload barrier")

    monkeypatch.setattr(storage, "upload", upload_then_wait)
    return entered, release, completed


def _fail_upload_number(
    storage: Any, monkeypatch: pytest.MonkeyPatch, failed_number: int
) -> list[int]:
    original = storage.upload
    calls = [0]

    def fail_selected(key: str, stream: Any, mime: str) -> None:
        calls[0] += 1
        if calls[0] == failed_number:
            raise OSError("controlled MinIO PUT failure")
        original(key, stream, mime)

    monkeypatch.setattr(storage, "upload", fail_selected)
    return calls


def _task_state(database: Any, task_id: int, plan_id: int) -> dict[str, Any]:
    with Session(database.engine) as session:
        task = session.get(GenerationTask, task_id)
        assert task is not None
        artifacts = list(
            session.execute(
                select(GeneratedFile, StorageObject)
                .join(StorageObject, StorageObject.id == GeneratedFile.storage_object_id)
                .where(GeneratedFile.task_id == task_id)
                .order_by(GeneratedFile.file_format)
            ).all()
        )
        return {
            "status": task.status,
            "phase": task.phase,
            "error_code": task.error_code,
            "file_id": task.file_id,
            "execution_owner": task.execution_owner,
            "lease_expires_at": task.execution_lease_expires_at,
            "attempt_count": task.execution_attempt_count,
            "exported_rows": task.exported_rows,
            "artifacts": [
                {
                    "file_id": int(generated.id),
                    "file_format": generated.file_format,
                    "file_name": generated.file_name,
                    "content_sha256": generated.content_sha256,
                    "size_bytes": generated.size_bytes,
                    "storage_id": int(stored.id),
                    "object_key": stored.object_key,
                    "object_type": stored.object_type,
                    "storage_size": stored.size_bytes,
                    "storage_sha256": stored.content_sha256,
                }
                for generated, stored in artifacts
            ],
            "storage_row_count": int(
                session.scalar(
                    select(func.count(StorageObject.id)).where(
                        StorageObject.object_type == "export"
                    )
                )
                or 0
            ),
            "plan_exists": session.get(DatabaseExportPlan, plan_id) is not None,
        }


def _assert_attempt_objects_absent(root_storage: Any, bucket: str, keys: list[str]) -> None:
    visible = {item.object_name for item in root_storage.list_objects(bucket, recursive=True)}
    assert not (set(keys) & visible)


def _artifact_bytes(root_storage: Any, bucket: str, key: str) -> bytes:
    response = root_storage.get_object(bucket, key)
    try:
        return response.read()
    finally:
        response.close()
        response.release_conn()


def _assert_published_artifacts(
    runtime: Any,
    database: Any,
    root_storage: Any,
    task_data: dict[str, Any],
    kind: str,
) -> dict[str, Any]:
    state = _task_state(database, task_data["task_id"], task_data["plan_id"])
    assert state["status"] == "succeeded"
    assert state["error_code"] is None
    assert state["lease_expires_at"] is None
    assert state["execution_owner"]
    assert state["plan_exists"] is False
    expected_formats = {"xlsx"} if kind == "database_export" else {"png", "pdf", "xlsx", "csv"}
    assert {item["file_format"] for item in state["artifacts"]} == expected_formats
    assert state["storage_row_count"] == len(expected_formats)
    assert state["file_id"] in {item["storage_id"] for item in state["artifacts"]}
    assert state["exported_rows"] == len(_ROWS)
    artifact_payloads: dict[str, bytes] = {}
    for item in state["artifacts"]:
        assert item["object_type"] == "export"
        assert item["size_bytes"] == item["storage_size"] > 0
        assert item["content_sha256"] == item["storage_sha256"]
        runtime.worker.generation.storage.verify(
            item["object_key"], item["size_bytes"], item["content_sha256"]
        )
        body = _artifact_bytes(
            root_storage, runtime.worker.generation.storage.bucket, item["object_key"]
        )
        artifact_payloads[item["file_format"]] = body
        if item["file_format"] == "xlsx":
            assert zipfile.is_zipfile(io.BytesIO(body))
            workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
            try:
                if kind == "database_export":
                    values = list(workbook["导出结果"].iter_rows(values_only=True))
                    assert values == [tuple(column["name"] for column in _COLUMNS)] + [
                        tuple(row) for row in _ROWS
                    ]
                else:
                    values = list(
                        workbook["Data"].iter_rows(min_row=2, max_col=3, values_only=True)
                    )
                    assert [(row[0], int(row[1]), row[2]) for row in values] == [
                        ("north", 12, "12"),
                        ("south", 7, "7"),
                        ("east", 4, "4"),
                    ]
            finally:
                workbook.close()
    if kind == "chart_render":
        png = Image.open(io.BytesIO(artifact_payloads["png"]))
        png.verify()
        png = Image.open(io.BytesIO(artifact_payloads["png"]))
        try:
            assert png.size == (1600, 900)
            assert png.info["chart_type"] == "bar"
            assert png.info["statistic"] == "sum"
        finally:
            png.close()

        pdf = fitz.open(stream=artifact_payloads["pdf"], filetype="pdf")
        try:
            assert len(pdf) == 1
            page = pdf[0]
            assert tuple(int(value) for value in (page.rect.width, page.rect.height)) == (1600, 900)
            images = page.get_images(full=True)
            assert len(images) == 1
            embedded = pdf.extract_image(images[0][0])
            pdf_chart = Image.open(io.BytesIO(embedded["image"]))
            png_chart = Image.open(io.BytesIO(artifact_payloads["png"]))
            try:
                assert pdf_chart.size == (1600, 900)
                difference = ImageChops.difference(
                    pdf_chart.convert("RGB"), png_chart.convert("RGB")
                )
                assert difference.getbbox() is None
            finally:
                pdf_chart.close()
                png_chart.close()
        finally:
            pdf.close()

        csv_rows = list(csv.DictReader(io.StringIO(artifact_payloads["csv"].decode("utf-8-sig"))))
        assert [(row["category"], row["amount"], row["metric_exact"]) for row in csv_rows] == [
            ("north", "12", "12"),
            ("south", "7", "7"),
            ("east", "4", "4"),
        ]
    return state


@pytest.mark.parametrize("kind", ["database_export", "chart_render"])
@pytest.mark.parametrize(
    ("stream_error", "initially_registered", "expected_error"),
    [
        (None, False, ErrorCode.QUERY_CONTEXT_EXPIRED.value),
        (ErrorCode.DATABASE_QUERY_TIMEOUT.value, True, ErrorCode.DATABASE_QUERY_TIMEOUT.value),
    ],
)
def test_missing_handle_or_remote_stream_error_fails_closed(
    runtime,
    database,
    ingestion_settings,
    monkeypatch,
    kind,
    stream_error,
    initially_registered,
    expected_error,
):
    task_data = _create_export_task(runtime, kind)
    worker = runtime.worker.generation
    keys = _track_attempt_keys(worker.storage, monkeypatch)
    with _controlled_database(
        runtime,
        task_data,
        stream_error=stream_error,
        initially_registered=initially_registered,
    ) as (_worker, transport):
        assert runtime.run(worker.handle_event, task_data["event_id"])
        state = _task_state(database, task_data["task_id"], task_data["plan_id"])
        assert state["status"] == "failed"
        assert state["error_code"] == expected_error
        assert state["file_id"] is None and state["artifacts"] == []
        assert state["storage_row_count"] == 0
        assert state["execution_owner"] and state["lease_expires_at"] is None
        assert state["attempt_count"] == 1 and state["plan_exists"] is False
        assert len(transport.stream_requests) == 1
        assert [item["handle"] for item in transport.release_requests] == [task_data["handle"]]
        assert transport.active_handles == set()
    _, root_storage, _ = ingestion_settings
    _assert_attempt_objects_absent(root_storage, worker.storage.bucket, keys)


@pytest.mark.parametrize("kind", ["database_export", "chart_render"])
@pytest.mark.parametrize(
    ("change", "expected_error"),
    [
        ("owner_disabled", ErrorCode.ACCOUNT_DISABLED.value),
        ("relation_revoked", ErrorCode.PERMISSION_DENIED.value),
    ],
)
def test_authorization_change_after_real_upload_prevents_publication(
    runtime, database, ingestion_settings, monkeypatch, kind, change, expected_error
):
    task_data = _create_export_task(runtime, kind)
    worker = runtime.worker.generation
    keys = _track_attempt_keys(worker.storage, monkeypatch)
    expected_uploads = 1 if kind == "database_export" else 4
    entered, release, completed = _pause_after_uploads(
        worker.storage, monkeypatch, expected_uploads
    )
    _, root_storage, _ = ingestion_settings
    with _controlled_database(runtime, task_data) as (_worker, transport):
        with ThreadPoolExecutor(max_workers=1) as pool:
            future = pool.submit(runtime.run, worker.handle_event, task_data["event_id"])
            try:
                assert entered.wait(timeout=30)
                assert completed[0] == expected_uploads == len(keys)
                for key in keys:
                    assert root_storage.stat_object(worker.storage.bucket, key).size > 0
                with Session(database.engine) as session, session.begin():
                    if change == "owner_disabled":
                        owner = session.get(User, task_data["owner_id"])
                        assert owner is not None
                        owner.status = "disabled"
                    else:
                        policy = session.get(RelationPolicy, task_data["policy_id"])
                        assert policy is not None
                        policy.enabled = False
                release.set()
            finally:
                release.set()
            assert future.result(timeout=30) is True

        state = _task_state(database, task_data["task_id"], task_data["plan_id"])
        assert state["status"] == "failed"
        assert state["error_code"] == expected_error
        assert state["file_id"] is None and state["artifacts"] == []
        assert state["storage_row_count"] == 0
        assert state["lease_expires_at"] is None and state["plan_exists"] is False
        assert transport.active_handles == set()
        assert len(transport.release_requests) == 1
    _assert_attempt_objects_absent(root_storage, worker.storage.bucket, keys)


@pytest.mark.parametrize("kind", ["database_export", "chart_render"])
def test_http_cancel_after_real_upload_discards_the_attempt(
    runtime, database, ingestion_settings, monkeypatch, kind
):
    task_data = _create_export_task(runtime, kind)
    worker = runtime.worker.generation
    keys = _track_attempt_keys(worker.storage, monkeypatch)
    expected_uploads = 1 if kind == "database_export" else 4
    entered, release, completed = _pause_after_uploads(
        worker.storage, monkeypatch, expected_uploads
    )
    _, root_storage, _ = ingestion_settings
    with _controlled_database(runtime, task_data) as (_worker, transport):
        with ThreadPoolExecutor(max_workers=1) as pool:
            future = pool.submit(runtime.run, worker.handle_event, task_data["event_id"])
            try:
                assert entered.wait(timeout=30)
                assert completed[0] == expected_uploads == len(keys)
                for key in keys:
                    assert root_storage.stat_object(worker.storage.bucket, key).size > 0
                cancelled = runtime.client.post(
                    f"/api/v1/generation-tasks/{task_data['task_id']}/cancel",
                    headers=runtime.admin,
                )
                assert cancelled.status_code == 200, cancelled.text
                assert cancelled.json()["data"]["status"] == "cancelled"
                assert cancelled.json()["data"]["file_id"] is None
                release.set()
            finally:
                release.set()
            assert future.result(timeout=30) is True

        state = _task_state(database, task_data["task_id"], task_data["plan_id"])
        assert state["status"] == "cancelled"
        assert state["file_id"] is None and state["artifacts"] == []
        assert state["storage_row_count"] == 0
        assert state["lease_expires_at"] is None and state["plan_exists"] is False
        assert transport.active_handles == set()
        assert len(transport.release_requests) == 1
    _assert_attempt_objects_absent(root_storage, worker.storage.bucket, keys)


def test_chart_partial_bundle_put_failure_removes_earlier_real_uploads(
    runtime, database, ingestion_settings, monkeypatch
):
    task_data = _create_export_task(runtime, "chart_render")
    worker = runtime.worker.generation
    keys = _track_attempt_keys(worker.storage, monkeypatch)
    calls = _fail_upload_number(worker.storage, monkeypatch, failed_number=2)
    _, root_storage, _ = ingestion_settings
    with _controlled_database(runtime, task_data) as (_worker, transport):
        assert runtime.run(worker.handle_event, task_data["event_id"])
        state = _task_state(database, task_data["task_id"], task_data["plan_id"])
        assert calls == [2]
        assert len(keys) == 2
        assert state["status"] == "failed"
        assert state["error_code"] == "GENERATION_FAILED"
        assert state["file_id"] is None and state["artifacts"] == []
        assert state["storage_row_count"] == 0
        assert state["lease_expires_at"] is None and state["plan_exists"] is False
        assert transport.active_handles == set()
        assert len(transport.release_requests) == 1
    _assert_attempt_objects_absent(root_storage, worker.storage.bucket, keys)


def _expire_task_lease(database: Any, task_id: int) -> None:
    with Session(database.engine) as session, session.begin():
        task = session.get(GenerationTask, task_id)
        assert task is not None and task.status == "running"
        task.execution_lease_expires_at = datetime.now(UTC) - timedelta(seconds=1)


@pytest.mark.parametrize("kind", ["database_export", "chart_render"])
def test_takeover_fences_old_export_attempt_while_new_owner_is_running(
    runtime, database, ingestion_settings, monkeypatch, kind
):
    task_data = _create_export_task(runtime, kind)
    worker = runtime.worker.generation
    worker.worker_config = WorkerConfig(lease_seconds=10, lease_renew_interval_seconds=1)
    keys = _track_attempt_keys(worker.storage, monkeypatch)
    old_entered, old_release = threading.Event(), threading.Event()
    new_entered, new_release = threading.Event(), threading.Event()
    barriers = ((old_entered, old_release), (new_entered, new_release))
    _, root_storage, _ = ingestion_settings
    with _controlled_database(runtime, task_data, stream_barriers=barriers) as (_worker, transport):
        with ThreadPoolExecutor(max_workers=2) as pool:
            old_future = pool.submit(runtime.run, worker.handle_event, task_data["event_id"])
            new_future = None
            try:
                assert old_entered.wait(timeout=30)
                old_claim = _task_state(database, task_data["task_id"], task_data["plan_id"])
                old_owner = old_claim["execution_owner"]
                assert old_claim["status"] == "running"
                assert old_claim["attempt_count"] == 1 and old_owner
                _expire_task_lease(database, task_data["task_id"])

                new_future = pool.submit(runtime.run, worker.handle_event, task_data["event_id"])
                assert new_entered.wait(timeout=30)
                active = _task_state(database, task_data["task_id"], task_data["plan_id"])
                new_owner = active["execution_owner"]
                assert active["status"] == "running"
                assert active["attempt_count"] == 2
                assert new_owner and new_owner != old_owner
                assert active["lease_expires_at"] > datetime.now(UTC)

                assert (
                    runtime.run(
                        worker._cleanup_export_plan,
                        task_data["task_id"],
                        task_data["handle"],
                        old_owner,
                    )
                    is None
                )
                active = _task_state(database, task_data["task_id"], task_data["plan_id"])
                assert active["status"] == "running" and active["execution_owner"] == new_owner
                assert active["plan_exists"] is True and active["artifacts"] == []
                assert active["storage_row_count"] == 0
                assert transport.active_handles == {task_data["handle"]}

                old_release.set()
                assert old_future.result(timeout=30) is False
                active = _task_state(database, task_data["task_id"], task_data["plan_id"])
                assert active["status"] == "running" and active["execution_owner"] == new_owner
                assert active["attempt_count"] == 2
                assert active["plan_exists"] is True and active["artifacts"] == []
                assert active["storage_row_count"] == 0
                assert keys == []
                assert transport.active_handles == {task_data["handle"]}

                new_release.set()
                assert new_future.result(timeout=45) is True
            finally:
                old_release.set()
                new_release.set()

        state = _assert_published_artifacts(runtime, database, root_storage, task_data, kind)
        assert state["attempt_count"] == 2
        assert state["execution_owner"] != old_owner
        assert transport.release_requests == [
            {
                "handle": task_data["handle"],
                "owner_user_id": task_data["owner_id"],
                "conversation_id": transport.stream_requests[1]["conversation_id"],
            }
        ]
        assert transport.active_handles == set()
        assert len(transport.stream_requests) == 2


@pytest.mark.parametrize("kind", ["database_export", "chart_render"])
def test_http_release_failure_retries_terminal_plan_cleanup_and_keeps_files(
    runtime, database, ingestion_settings, kind
):
    task_data = _create_export_task(runtime, kind)
    worker = runtime.worker.generation
    _, root_storage, _ = ingestion_settings
    with _controlled_database(runtime, task_data, release_http_failures=1) as (_worker, transport):
        assert runtime.run(worker.handle_event, task_data["event_id"]) is False
        first = _assert_published_artifacts(runtime, database, root_storage, task_data, kind)
        assert transport.release_requests == [
            {
                "handle": task_data["handle"],
                "owner_user_id": task_data["owner_id"],
                "conversation_id": transport.stream_requests[0]["conversation_id"],
            }
        ]
        assert transport.active_handles == {task_data["handle"]}
        assert first["artifacts"]

        assert runtime.run(worker.handle_event, task_data["event_id"]) is True
        final = _assert_published_artifacts(runtime, database, root_storage, task_data, kind)
        assert final["artifacts"] == first["artifacts"]
        assert len(transport.release_requests) == 2
        assert transport.active_handles == set()


@pytest.mark.parametrize("kind", ["database_export", "chart_render"])
def test_release_connect_error_uses_owner_ttl_fallback_without_replay(
    runtime, database, ingestion_settings, kind
):
    task_data = _create_export_task(runtime, kind)
    worker = runtime.worker.generation
    _, root_storage, _ = ingestion_settings
    with _controlled_database(runtime, task_data, release_connect_failure=True) as (
        _worker,
        transport,
    ):
        assert runtime.run(worker.handle_event, task_data["event_id"]) is True
        state = _assert_published_artifacts(runtime, database, root_storage, task_data, kind)
        assert state["artifacts"]
        assert len(transport.release_requests) == 1
        assert transport.active_handles == {task_data["handle"]}
