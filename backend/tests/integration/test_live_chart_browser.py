"""Opt-in synthetic Qwen → browser → independent Worker/MCP acceptance."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import sys
import time
import zipfile
from decimal import Decimal
from pathlib import Path

import httpx
import pytest
from conftest import ROOT
from openpyxl import load_workbook
from phase8_helpers import free_port, refresh_schema, stop_process
from PIL import Image
from psycopg import sql
from sqlalchemy.orm import Session
from test_phase8_live import _build_live_runtime, _post

from enterprise_mind.agent.generation_models import GeneratedFile
from enterprise_mind.ingestion.models import StorageObject

pytest_plugins = ("phase8_helpers",)
pytestmark = [
    pytest.mark.integration,
    pytest.mark.skipif(
        os.getenv("EM_TEST_CHART_BROWSER") != "1" or os.getenv("EM_TEST_LIVE_MODELS") != "1",
        reason="Explicit live-model and chart-browser opt-in required",
    ),
]


@pytest.fixture
def chart_live_runtime(database, ingestion_settings, phase8_mcp, tmp_path):
    services, _, _ = ingestion_settings
    yield from _build_live_runtime(
        database, phase8_mcp, tmp_path, services=services, require_rag=True
    )


def test_three_chart_types_live_agent_worker_and_two_browsers(
    chart_live_runtime, ingestion_settings
):
    runtime = chart_live_runtime
    client = runtime["client"]
    admin = runtime["admin"]
    business = runtime["business"]
    services, root_storage, _ = ingestion_settings
    output = Path(os.environ["EM_CHART_BROWSER_OUTPUT"])
    output.mkdir(parents=True, exist_ok=True)
    with business.admin() as connection:
        connection.execute(
            sql.SQL(
                "UPDATE {}.orders SET department = CASE id % 3 "
                "WHEN 0 THEN 'sales' WHEN 1 THEN 'support' ELSE 'finance' END"
            ).format(sql.Identifier(business.sales_schema))
        )
        connection.execute(
            sql.SQL(
                "INSERT INTO {}.orders (id,order_date,amount,status,department) "
                "VALUES (236,'2026-01-01',1000,'unpaid','sales')"
            ).format(sql.Identifier(business.sales_schema))
        )
    response = client.put(
        "/api/v1/admin/database-mcp/source",
        headers={**admin, "Idempotency-Key": "chart-browser-source"},
        json=business.source_body,
    )
    assert response.status_code == 201
    _post(client, admin, "/api/v1/admin/database-mcp/schema-refresh", {}, status=202)
    assert (
        refresh_schema(
            runtime["database"].settings, runtime["mcp"].endpoint, business.token, limit=1
        )
        == 1
    )
    response = client.put(
        "/api/v1/admin/database-mcp/relation-policies",
        headers=admin,
        json={
            "schema_name": business.sales_schema,
            "relation_name": business.orders_table,
            "enabled": True,
            "department_ids": [],
        },
    )
    assert response.status_code == 200
    children = []
    logs = []
    trace = output / "worker-events.jsonl"
    worker_code = (
        "import asyncio,json,os; from pathlib import Path; "
        "from enterprise_mind.generation.worker import GenerationWorker; "
        "from enterprise_mind.ingestion.runtime import run_processing; "
        "from enterprise_mind.infrastructure.settings import "
        "DatabaseSettings,load_application_config\n"
        "original=GenerationWorker.handle_event\n"
        "async def observed(self,event_id):\n"
        " result=await original(self,event_id)\n"
        " with Path(os.environ['EM_CHART_WORKER_TRACE']).open('a') as f: "
        "f.write(json.dumps({'pid':os.getpid(),'ppid':os.getppid(),"
        "'event_id':event_id,'handled':bool(result)})+'\\n')\n"
        " return result\n"
        "GenerationWorker.handle_event=observed\n"
        "s=DatabaseSettings(); asyncio.run(run_processing('worker',s,"
        "load_application_config(s.application_config),once=False))"
    )
    env = {**runtime["env"], "EM_CHART_WORKER_TRACE": str(trace)}
    node = shutil.which("node")
    assert node is not None
    port = free_port()
    base_url = f"http://127.0.0.1:{port}"
    try:
        for name, arguments, child_env, cwd in (
            (
                "dispatcher",
                [sys.executable, "-m", "enterprise_mind.cli", "dispatch"],
                env,
                ROOT,
            ),
            ("worker", [sys.executable, "-c", worker_code], env, ROOT),
            (
                "vite",
                [node, "node_modules/vite/bin/vite.js", "--host", "127.0.0.1", "--port", str(port)],
                {**os.environ, "P10_API_PROXY": str(client.base_url).rstrip("/")},
                ROOT / "frontend",
            ),
        ):
            log = (output / f"{name}.log").open("wb")
            logs.append(log)
            process = subprocess.Popen(
                arguments,
                cwd=cwd,
                env=child_env,
                stdout=log,
                stderr=subprocess.STDOUT,
                creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0,
            )
            children.append(process)
        deadline = time.monotonic() + 40
        with httpx.Client(trust_env=False) as browser_client:
            while time.monotonic() < deadline:
                assert all(child.poll() is None for child in children), "child_start_failed"
                try:
                    if browser_client.get(base_url).status_code == 200:
                        break
                except httpx.TransportError:
                    pass
                time.sleep(0.2)
            else:
                pytest.fail("vite_not_ready")
        charts = []
        requested_types = set(os.getenv("EM_CHART_TYPES", "bar,line,pie").split(","))
        for kind, label, dimension, points in (
            ("bar", "柱状图", "department", 3),
            ("line", "折线图", "order_date", 28),
            ("pie", "饼图", "department", 3),
        ):
            if kind not in requested_types:
                continue
            charts.append(
                {
                    "type": kind,
                    "points": points,
                    "question": (
                        f"请查询授权数据库 {business.orders_relation} "
                        "中 status 为 paid 的全部订单，"
                        f"按 {dimension} 分组汇总 SUM(amount)，"
                        f"维度列为 {dimension}，指标列为 total_amount，"
                        f"按 {dimension} 升序，生成全部聚合结果的{label} PNG 预览。"
                    ),
                }
            )
        result = subprocess.run(
            [node, str(ROOT / "frontend/tests/live-chart-browser.mjs")],
            input=json.dumps(
                {
                    "baseUrl": base_url,
                    "username": "初始管理员",
                    "password": runtime["database"].admin_password,
                    "charts": charts,
                    "output": str(output),
                }
            ),
            text=True,
            capture_output=True,
            timeout=2400,
            creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0,
        )
        assert result.returncode == 0, "browser_probe_failed; see sanitized browser-result.json"
        report = json.loads((output / "browser-result.json").read_text(encoding="utf-8"))
        assert report["status"] == "passed"
        assert len(report["cases"]) == len(charts)
        for case in report["cases"]:
            assert len(case["downloads"]) == 8
            for download in case["downloads"]:
                saved = Path(download["saved"]).read_bytes()
                assert hashlib.sha256(saved).hexdigest() == download["sha256"]
                assert len(saved) == download["bytes"]
                with Session(runtime["database"].engine) as session:
                    file = session.get(GeneratedFile, int(download["fileId"]))
                    assert file is not None and file.task_id == int(case["taskId"])
                    storage = session.get(StorageObject, file.storage_object_id)
                    assert storage is not None
                    stored = root_storage.get_object(storage.bucket_name, storage.object_key)
                    try:
                        payload = stored.read()
                    finally:
                        stored.close()
                        stored.release_conn()
                    assert (
                        hashlib.sha256(payload).hexdigest()
                        == file.content_sha256
                        == download["sha256"]
                    )
                    assert len(payload) == file.size_bytes == download["bytes"]
                if download["format"] == "csv":
                    rows = list(csv.DictReader(io.StringIO(payload.decode("utf-8-sig"))))
                    assert len(rows) == case["points"]
                    assert sum(Decimal(row["metric_exact"]) for row in rows) == Decimal("2350.00")
                    if case["type"] == "line":
                        expected = {
                            f"2026-01-{day:02d}": Decimal(
                                sum(1 for i in range(1, 236) if i % 28 + 1 == day)
                            )
                            * 10
                            for day in range(1, 29)
                        }
                        assert {
                            row["order_date"]: Decimal(row["metric_exact"]) for row in rows
                        } == expected
                    else:
                        assert {
                            row["department"]: Decimal(row["metric_exact"]) for row in rows
                        } == {
                            "sales": Decimal("780"),
                            "support": Decimal("790"),
                            "finance": Decimal("780"),
                        }
                    assert all(
                        row["__source_refs"]
                        and row["__queried_at"]
                        and re.search(r"\bSUM\s*\(", row["__statistic"], re.I)
                        and "amount" in row["__statistic"]
                        for row in rows
                    )
                elif download["format"] == "png":
                    with Image.open(io.BytesIO(payload)) as image:
                        assert image.size == (1600, 900)
                elif download["format"] == "pdf":
                    assert payload.startswith(b"%PDF-")
                elif download["format"] == "xlsx":
                    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                        assert any(
                            name.startswith("xl/charts/chart") for name in archive.namelist()
                        )
                    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
                    try:
                        assert workbook.sheetnames == ["Chart", "Data", "Source"]
                    finally:
                        workbook.close()
        events = [json.loads(line) for line in trace.read_text().splitlines()]
        assert sum(item["handled"] for item in events) >= len(charts)
        # The Windows venv launcher may own a base-CPython child. Require the
        # trace to come from the exact worker or its directly verified child.
        assert len({item["pid"] for item in events}) == 1
        assert all(
            item["pid"] == children[1].pid or item["ppid"] == children[1].pid for item in events
        )
        runtime["record"](
            "chart_agent_browser_worker",
            "passed",
            chart_types=len(charts),
            browser_count=2,
            saved_files=len(charts) * 8,
            source_and_exact_totals=True,
        )
    finally:
        for child in reversed(children):
            stop_process(child)
        for log in logs:
            log.close()
        (output / "process-cleanup.json").write_text(
            json.dumps(
                {
                    "children_exited": all(child.poll() is not None for child in children),
                    "pids": [child.pid for child in children],
                    "random_bucket": services.minio_bucket,
                    "random_stream": services.redis_stream,
                }
            ),
            encoding="utf-8",
        )
