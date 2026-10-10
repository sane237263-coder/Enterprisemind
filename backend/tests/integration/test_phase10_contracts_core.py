# ruff: noqa: F811

"""Phase 10 direct contracts for Memory, draft freezing, and generated imports.

These tests intentionally use the application session factory and HTTP routes
instead of SQLite or an in-memory repository.  They are integration tests and
are skipped by the shared fixture unless the isolated PostgreSQL (and, for
generated imports, Redis/MinIO) environment is explicitly configured.
"""

from __future__ import annotations

import asyncio
import hashlib
import io
import uuid
from datetime import UTC, datetime, timedelta
from typing import Any

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session
from test_agent import agent as agent_fixture  # noqa: F401
from test_agent import new_run
from test_ingestion import ingestion_runtime  # noqa: F401
from test_management import auth

from enterprise_mind.agent.generation_models import ContentDraft, GeneratedFile, GenerationTask
from enterprise_mind.agent.models import Conversation, Message
from enterprise_mind.api.errors import APIError, ErrorCode
from enterprise_mind.app import create_app
from enterprise_mind.generation import service as generation_service
from enterprise_mind.identity.models import User
from enterprise_mind.identity.service import CurrentIdentity
from enterprise_mind.ingestion.models import (
    Document,
    OutboxEvent,
    Revision,
    StorageObject,
)
from enterprise_mind.knowledge.pending_models import PendingAction
from enterprise_mind.memory import service as memory_service
from enterprise_mind.memory.extraction import MemoryCandidateSet
from enterprise_mind.memory.models import Memory, MemoryState

pytestmark = pytest.mark.integration


def _post(client: TestClient, headers: dict[str, str], path: str, body: dict[str, Any], key: str):
    return client.post(path, headers={**headers, "Idempotency-Key": key}, json=body)


def _create_user(client: TestClient, admin: dict[str, str], username: str) -> dict[str, Any]:
    response = _post(
        client,
        admin,
        "/api/v1/admin/users",
        {
            "username": username,
            "password": "Phase10-Memory-Contract-2026!",
            "role": "admin",
        },
        f"create-{username}",
    )
    assert response.status_code == 201, response.text
    return response.json()["data"]


def test_memory_api_is_owner_scoped_and_context_has_no_cross_user_rows(database):
    database.initialize()
    app = create_app(database.settings)
    with TestClient(app) as client:
        admin = auth(client, "初始管理员", database.admin_password)
        user_a = _create_user(client, admin, "phase10-memory-a")
        user_b = _create_user(client, admin, "phase10-memory-b")
        token_a = auth(client, "phase10-memory-a", "Phase10-Memory-Contract-2026!")
        token_b = auth(client, "phase10-memory-b", "Phase10-Memory-Contract-2026!")

        created = client.post(
            "/api/v1/me/memories",
            headers={**token_a, "Idempotency-Key": "memory-owner-create"},
            json={
                "category": "preference",
                "subject_key": "answer-style",
                "content": "先给结论，再给证据。",
            },
        )
        assert created.status_code == 201, created.text
        memory_id = created.json()["data"]["id"]
        assert client.get("/api/v1/me/memories", headers=token_b).json()["data"] == []
        # The update/delete routes are owner scoped even when the identifier is valid.
        denied_update = client.patch(
            f"/api/v1/me/memories/{memory_id}",
            headers=token_b,
            json={"expected_revision": 1, "content": "跨用户修改不应成功。"},
        )
        assert denied_update.status_code == 404, denied_update.text
        denied_delete = client.request(
            "DELETE",
            f"/api/v1/me/memories/{memory_id}",
            headers=token_b,
            json={"expected_revision": 1},
        )
        assert denied_delete.status_code == 404, denied_delete.text

        async def context_rows() -> tuple[list[int], list[int]]:
            async with app.state.sessions() as session:
                a_rows = await memory_service.select_for_context(
                    session, int(user_a["id"]), "answer-style"
                )
                b_rows = await memory_service.select_for_context(
                    session, int(user_b["id"]), "answer-style"
                )
                return [int(row.user_id) for row in a_rows], [int(row.user_id) for row in b_rows]

        a_ids, b_ids = client.portal.call(context_rows)
        assert a_ids == [int(user_a["id"])]
        assert b_ids == []


def test_memory_read_failure_degrades_ordinary_qna(agent_fixture, monkeypatch):
    client, headers, manager = agent_fixture

    async def unavailable(*args, **kwargs):
        raise RuntimeError("memory store unavailable")

    monkeypatch.setattr(memory_service, "select_for_context", unavailable)
    conversation_id, run_id = new_run(agent_fixture, "请写一段不依赖 Memory 的欢迎词。")
    assert client.portal.call(manager.execute, run_id)
    messages = client.get(f"/api/v1/conversations/{conversation_id}/messages", headers=headers)
    assert messages.status_code == 200, messages.text
    assert messages.json()["data"][-1]["status"] == "completed"


@pytest.mark.parametrize(
    ("sensitive_text", "secret_fragment"),
    [
        ("我的 password=fixture-password，请记住先给结论。", "fixture-password"),
        ("我的 token=fixture-token，请记住先给结论。", "fixture-token"),
        ("我的 api_key=fixture-api-key，请记住先给结论。", "fixture-api-key"),
        ("我的 secret=fixture-secret，请记住先给结论。", "fixture-secret"),
        (
            "连接 postgres://fixture:fixture-pass@localhost/business，请记住先给结论。",
            "fixture-pass",
        ),
        (
            "-----BEGIN PRIVATE KEY-----\nfixture-private-key\n"
            "-----END PRIVATE KEY-----，请记住先给结论。",
            "fixture-private-key",
        ),
    ],
)
def test_memory_extractor_receives_sanitized_text_before_model_call(
    agent_fixture, database, sensitive_text: str, secret_fragment: str
):
    client, _, manager = agent_fixture
    with Session(database.engine) as session:
        user_id = database.user_id()
        conversation = Conversation(owner_user_id=user_id, title="phase10-secret-memory")
        session.add(conversation)
        session.flush()
        message = Message(
            conversation_id=conversation.id,
            sequence=1,
            role="user",
            content_text=sensitive_text,
            status="completed",
        )
        session.add(message)
        session.flush()
        message_id = message.id
        session.commit()
    seen: list[str] = []

    class ExtractorModel:
        async def extract_memory(self, content: str) -> MemoryCandidateSet:
            seen.append(content)
            return MemoryCandidateSet()

        async def close(self):
            pass

    manager.model = ExtractorModel()
    # Call the same best-effort extraction entry used after a completed run.
    client.portal.call(manager._extract_memory, user_id, message_id, 0)
    assert seen and secret_fragment not in seen[0] and "先给结论" in seen[0]


def test_memory_revision_race_and_stale_automatic_candidate_are_durable(database):
    database.initialize()
    app = create_app(database.settings)
    with TestClient(app) as client:
        admin = auth(client, "初始管理员", database.admin_password)
        user = _create_user(client, admin, "phase10-memory-race")
        token = auth(client, "phase10-memory-race", "Phase10-Memory-Contract-2026!")
        created = client.post(
            "/api/v1/me/memories",
            headers={**token, "Idempotency-Key": "memory-race-create"},
            json={
                "category": "preference",
                "subject_key": "report-style",
                "content": "先给摘要。",
            },
        )
        assert created.status_code == 201, created.text
        memory_id = int(created.json()["data"]["id"])

        async def run_race() -> list[tuple[str, int | str]]:
            barrier = asyncio.Barrier(2)

            async def attempt(content: str) -> tuple[str, int | str]:
                async with app.state.sessions() as session:
                    try:
                        async with session.begin():
                            await barrier.wait()
                            item, _, _ = await memory_service.save_for_user(
                                session,
                                int(user["id"]),
                                memory_id=memory_id,
                                expected_revision=1,
                                content=content,
                                category="preference",
                                subject_key="report-style",
                                origin="manual",
                            )
                            return "ok", item.revision
                    except APIError as error:
                        return "error", error.code.value

            return list(
                await asyncio.gather(
                    attempt("版本一的正文。"),
                    attempt("版本二的正文。"),
                )
            )

        results = client.portal.call(run_race)
        assert sorted(results) == [("error", ErrorCode.ROW_VERSION_CONFLICT.value), ("ok", 2)]

        async def apply_stale_candidate() -> dict[str, Any]:
            async with app.state.sessions() as session:
                async with session.begin():
                    epoch = await memory_service.current_epoch(session, int(user["id"]))
            async with app.state.sessions() as session:
                async with session.begin():
                    deleted_epoch = await memory_service.delete_for_user(
                        session, int(user["id"]), memory_id, expected_revision=2
                    )
            async with app.state.sessions() as session:
                async with session.begin():
                    result = await memory_service.apply_candidates(
                        session,
                        int(user["id"]),
                        epoch,
                        source_message_id=1,
                        candidates=[
                            {
                                "category": "preference",
                                "subject_key": "report-style",
                                "content": "旧自动提取不应复活。",
                            }
                        ],
                    )
                    result["deleted_epoch"] = deleted_epoch
                    return result

        stale = client.portal.call(apply_stale_candidate)
        assert stale["status"] == "stale"
        assert stale["deleted_epoch"] == stale["epoch"]
        with Session(database.engine) as session:
            assert session.scalar(select(Memory).where(Memory.id == memory_id)) is None
            state = session.get(MemoryState, int(user["id"]))
            assert state is not None and state.memory_epoch == stale["epoch"]


def test_duplicate_automatic_source_message_is_idempotent(database):
    database.initialize()
    app = create_app(database.settings)
    with TestClient(app) as client:
        auth(client, "初始管理员", database.admin_password)

        async def create_source() -> tuple[int, int, int, dict[str, Any]]:
            async with app.state.sessions() as session, session.begin():
                owner = await session.scalar(select(User).where(User.username == "初始管理员"))
                assert owner is not None
                conversation = Conversation(owner_user_id=owner.id, title="phase10-memory-source")
                session.add(conversation)
                await session.flush()
                message = Message(
                    conversation_id=conversation.id,
                    sequence=1,
                    role="user",
                    content_text="请记住我的输出偏好。",
                    status="completed",
                )
                session.add(message)
                await session.flush()
                captured = await memory_service.current_epoch(session, owner.id)
                candidate = {
                    "category": "preference",
                    "subject_key": "source-idempotency",
                    "content": "重复来源消息只保存一次。",
                }
                return owner.id, message.id, captured, candidate

        owner_id, message_id, captured_epoch, candidate = client.portal.call(create_source)

        async def apply_once(epoch: int) -> dict[str, Any]:
            async with app.state.sessions() as session, session.begin():
                return await memory_service.apply_candidates(
                    session, owner_id, epoch, message_id, [candidate]
                )

        first = client.portal.call(apply_once, captured_epoch)
        second = client.portal.call(apply_once, first["epoch"])
        assert first["status"] == "applied" and first["saved"] == 1
        assert second["status"] == "applied" and second["saved"] == 0
        with Session(database.engine) as session:
            rows = list(
                session.scalars(
                    select(Memory).where(
                        Memory.user_id == owner_id,
                        Memory.subject_key == "source-idempotency",
                    )
                )
            )
            assert len(rows) == 1
            assert rows[0].source_message_id == message_id
            assert rows[0].revision == 1


def test_manual_memory_update_wins_over_old_automatic_candidate(database):
    database.initialize()
    app = create_app(database.settings)
    with TestClient(app) as client:
        token = auth(client, "初始管理员", database.admin_password)
        created = client.post(
            "/api/v1/me/memories",
            headers={**token, "Idempotency-Key": "memory-manual-priority-create"},
            json={
                "category": "preference",
                "subject_key": "manual-priority",
                "content": "用户初始偏好。",
            },
        )
        assert created.status_code == 201, created.text
        memory_id = int(created.json()["data"]["id"])
        initial_epoch = int(created.json()["data"]["memory_epoch"])

        updated = client.patch(
            f"/api/v1/me/memories/{memory_id}",
            headers=token,
            json={
                "expected_revision": 1,
                "content": "用户手动确认的新偏好。",
            },
        )
        assert updated.status_code == 200, updated.text
        assert updated.json()["data"]["revision"] == 2
        manual_epoch = int(updated.json()["data"]["memory_epoch"])
        assert manual_epoch > initial_epoch

        async def owner_id() -> int:
            async with app.state.sessions() as session:
                owner = await session.scalar(select(User.id).where(User.username == "初始管理员"))
                assert owner is not None
                return int(owner)

        resolved_owner = client.portal.call(owner_id)

        async def create_source_message() -> int:
            async with app.state.sessions() as session, session.begin():
                conversation = Conversation(
                    owner_user_id=resolved_owner, title="phase10-memory-manual-priority"
                )
                session.add(conversation)
                await session.flush()
                message = Message(
                    conversation_id=conversation.id,
                    sequence=1,
                    role="user",
                    content_text="请记住我的偏好。",
                    status="completed",
                )
                session.add(message)
                await session.flush()
                return int(message.id)

        source_message_id = client.portal.call(create_source_message)

        async def apply_old() -> dict[str, Any]:
            async with app.state.sessions() as session, session.begin():
                return await memory_service.apply_candidates(
                    session,
                    resolved_owner,
                    initial_epoch,
                    source_message_id=source_message_id,
                    candidates=[
                        {
                            "category": "preference",
                            "subject_key": "manual-priority",
                            "content": "旧自动提取值。",
                        }
                    ],
                )

        async def apply_current() -> dict[str, Any]:
            async with app.state.sessions() as session, session.begin():
                return await memory_service.apply_candidates(
                    session,
                    resolved_owner,
                    manual_epoch,
                    source_message_id=source_message_id,
                    candidates=[
                        {
                            "category": "preference",
                            "subject_key": "manual-priority",
                            "content": "旧自动提取值。",
                        }
                    ],
                )

        stale = client.portal.call(apply_old)
        current = client.portal.call(apply_current)
        assert stale == {"status": "stale", "saved": 0, "epoch": manual_epoch}
        assert current == {"status": "applied", "saved": 0, "epoch": manual_epoch}
        with Session(database.engine) as session:
            row = session.get(Memory, memory_id)
            assert row is not None
            assert row.content == "用户手动确认的新偏好。"
            assert row.origin == "manual"
            assert row.revision == 2


def test_draft_task_freezes_body_tables_sources_and_revision(database):
    database.initialize()
    app = create_app(database.settings)
    with TestClient(app) as client:
        auth(client, "初始管理员", database.admin_password)

        async def create_and_freeze() -> tuple[int, int, dict[str, Any], dict[str, Any]]:
            async with app.state.sessions() as session, session.begin():
                user = await session.scalar(select(User).where(User.username == "初始管理员"))
                assert user is not None
                identity = CurrentIdentity(user, None, user.credential_version)
                # This public document is a real authorized source row, allowing
                # the service to perform its normal source revalidation.
                source_object = StorageObject(
                    bucket_name="phase10-contract",
                    object_key="original/source.md",
                    object_type="original",
                    size_bytes=1,
                    content_type="text/markdown",
                    content_sha256="a" * 64,
                    status="available",
                )
                session.add(source_object)
                await session.flush()
                from enterprise_mind.knowledge.models import KnowledgeBase

                base = KnowledgeBase(
                    name="Phase10 Source", name_key="phase10-source", type="public"
                )
                session.add(base)
                await session.flush()
                document = Document(
                    knowledge_base_id=base.id,
                    name="source.md",
                    name_key="source.md",
                    created_by=user.id,
                    processing_status="success",
                    latest_revision_no=1,
                )
                session.add(document)
                await session.flush()
                revision = Revision(
                    document_id=document.id,
                    revision_no=1,
                    trigger_type="initial_upload",
                    original_object_id=source_object.id,
                    original_file_name="source.md",
                    file_extension="md",
                    mime_type="text/markdown",
                    size_bytes=1,
                    content_sha256="a" * 64,
                    processing_status="success",
                    initiated_by=user.id,
                )
                session.add(revision)
                await session.flush()
                document.active_revision_id = revision.id
                source_refs = [
                    {
                        "source_type": "knowledge",
                        "document_id": str(document.id),
                        "revision_id": str(revision.id),
                        "title": "source.md",
                    }
                ]
                draft = await generation_service.create_draft(
                    session,
                    identity,
                    title="D1 标题",
                    body_markdown="# D1 正文\n\n冻结前的正文。",
                    tables=[
                        {
                            "title": "D1 表格",
                            "columns": ["字段", "值"],
                            "rows": [["A", "1"]],
                        }
                    ],
                    source_refs=source_refs,
                )
                task = await generation_service.create_render_task(
                    session,
                    identity,
                    draft_id=draft.id,
                    expected_revision=draft.revision,
                    file_format="md",
                    file_name="d1.md",
                    idempotency_key="phase10-draft-render",
                )
                frozen = dict(task.input_ref)
                await generation_service.update_draft(
                    session,
                    identity,
                    draft.id,
                    expected_revision=draft.revision,
                    title="D2 标题",
                    body_markdown="# D2 正文\n\n后续编辑。",
                    tables=[
                        {
                            "title": "D2 表格",
                            "columns": ["字段", "值"],
                            "rows": [["B", "2"]],
                        }
                    ],
                    source_refs=[],
                )
                return draft.id, task.id, frozen, source_refs

        draft_id, task_id, frozen, source_refs = client.portal.call(create_and_freeze)
        assert frozen["title"] == "D1 标题"
        assert frozen["body_markdown"].startswith("# D1")
        assert frozen["tables"] == [
            {"title": "D1 表格", "columns": ["字段", "值"], "rows": [["A", "1"]]}
        ]
        assert frozen["source_refs"] == source_refs
        assert frozen["draft_revision"] == 1
        with Session(database.engine) as session:
            task = session.get(GenerationTask, task_id)
            draft = session.get(ContentDraft, draft_id)
            assert task is not None and draft is not None
            assert task.input_ref["body_markdown"].startswith("# D1")
            assert task.source_refs == source_refs
            assert draft.revision == 2 and draft.body_markdown.startswith("# D2")


def test_frozen_draft_executes_d1_in_all_formats_before_d2(runtime, database, tmp_path):
    async def create_tasks() -> tuple[int, list[tuple[int, str]], dict[str, Any]]:
        async with runtime.worker.sessions() as session, session.begin():
            user = await session.scalar(select(User).where(User.username == "初始管理员"))
            assert user is not None
            identity = CurrentIdentity(user, None, user.credential_version)
            source_object = StorageObject(
                bucket_name="phase10-contract-render",
                object_key="original/render-source.md",
                object_type="original",
                size_bytes=1,
                content_type="text/markdown",
                content_sha256="c" * 64,
                status="available",
            )
            session.add(source_object)
            await session.flush()
            from enterprise_mind.knowledge.models import KnowledgeBase

            base = KnowledgeBase(
                name="Phase10 Render Source",
                name_key="phase10-render-source",
                type="public",
            )
            session.add(base)
            await session.flush()
            document = Document(
                knowledge_base_id=base.id,
                name="render-source.md",
                name_key="render-source.md",
                created_by=user.id,
                processing_status="success",
                latest_revision_no=1,
            )
            session.add(document)
            await session.flush()
            revision = Revision(
                document_id=document.id,
                revision_no=1,
                trigger_type="initial_upload",
                original_object_id=source_object.id,
                original_file_name="render-source.md",
                file_extension="md",
                mime_type="text/markdown",
                size_bytes=1,
                content_sha256="c" * 64,
                processing_status="success",
                initiated_by=user.id,
            )
            session.add(revision)
            await session.flush()
            document.active_revision_id = revision.id
            source_refs = [
                {
                    "source_type": "knowledge",
                    "document_id": str(document.id),
                    "revision_id": str(revision.id),
                    "title": "render-source.md",
                }
            ]
            draft = await generation_service.create_draft(
                session,
                identity,
                title="D1 标题",
                body_markdown="# D1 正文\n\n冻结正文不能被 D2 替换。",
                tables=[
                    {
                        "title": "D1 表格",
                        "columns": ["字段", "值"],
                        "rows": [["A", "1"]],
                    }
                ],
                source_refs=source_refs,
            )
            tasks: list[tuple[int, str]] = []
            for file_format in ("md", "docx", "pdf", "xlsx"):
                task = await generation_service.create_render_task(
                    session,
                    identity,
                    draft_id=draft.id,
                    expected_revision=draft.revision,
                    file_format=file_format,
                    file_name=f"phase10-d1.{file_format}",
                    idempotency_key=f"phase10-d1-{file_format}",
                )
                tasks.append((task.id, file_format))
            revoked_task = await generation_service.create_render_task(
                session,
                identity,
                draft_id=draft.id,
                expected_revision=draft.revision,
                file_format="md",
                file_name="phase10-d1-revoked.md",
                idempotency_key="phase10-d1-revoked",
            )
            frozen = dict((await session.get(GenerationTask, tasks[0][0])).input_ref)
            await generation_service.update_draft(
                session,
                identity,
                draft.id,
                expected_revision=draft.revision,
                title="D2 标题",
                body_markdown="# D2 正文\n\n修改后的正文。",
                tables=[
                    {
                        "title": "D2 表格",
                        "columns": ["字段", "值"],
                        "rows": [["B", "2"]],
                    }
                ],
                source_refs=[],
            )
            return draft.id, [*tasks, (revoked_task.id, "md")], frozen

    draft_id, task_formats_with_revoked, frozen = runtime.run(create_tasks)
    task_formats = task_formats_with_revoked[:4]
    revoked_task_id = task_formats_with_revoked[4][0]
    assert frozen["body_markdown"].startswith("# D1")
    assert frozen["tables"] == [
        {"title": "D1 表格", "columns": ["字段", "值"], "rows": [["A", "1"]]}
    ]
    assert frozen["draft_revision"] == 1

    async def revoke_source() -> None:
        async with runtime.worker.sessions() as session, session.begin():
            user = await session.scalar(
                select(User).where(User.username == "初始管理员").with_for_update()
            )
            assert user is not None
            user.status = "disabled"

    async def restore_source() -> None:
        async with runtime.worker.sessions() as session, session.begin():
            user = await session.scalar(
                select(User).where(User.username == "初始管理员").with_for_update()
            )
            assert user is not None
            user.status = "active"

    runtime.run(revoke_source)
    with Session(database.engine) as session:
        revoked_event = session.scalar(
            select(OutboxEvent).where(
                OutboxEvent.aggregate_type == "generation_task",
                OutboxEvent.aggregate_id == revoked_task_id,
                OutboxEvent.event_type == "generation.requested",
            )
        )
        assert revoked_event is not None
        revoked_event_id = revoked_event.id
    assert runtime.run(runtime.worker.generation.handle_event, revoked_event_id)
    with Session(database.engine) as session:
        revoked = session.get(GenerationTask, revoked_task_id)
        assert revoked is not None
        assert revoked.status == "failed"
        assert revoked.error_code == ErrorCode.ACCOUNT_DISABLED.value
        assert (
            session.scalar(select(GeneratedFile.id).where(GeneratedFile.task_id == revoked_task_id))
            is None
        )
    runtime.run(restore_source)

    async def reject_old_revision() -> None:
        async with runtime.worker.sessions() as session, session.begin():
            user = await session.scalar(select(User).where(User.username == "初始管理员"))
            assert user is not None
            identity = CurrentIdentity(user, None, user.credential_version)
            with pytest.raises(APIError) as error:
                await generation_service.create_render_task(
                    session,
                    identity,
                    draft_id=draft_id,
                    expected_revision=1,
                    file_format="md",
                    file_name="old.md",
                    idempotency_key="phase10-old-revision",
                )
            assert error.value.code == ErrorCode.ROW_VERSION_CONFLICT

    runtime.run(reject_old_revision)

    async def create_d2_task() -> tuple[int, dict[str, Any]]:
        async with runtime.worker.sessions() as session, session.begin():
            user = await session.scalar(select(User).where(User.username == "初始管理员"))
            assert user is not None
            identity = CurrentIdentity(user, None, user.credential_version)
            task = await generation_service.create_render_task(
                session,
                identity,
                draft_id=draft_id,
                expected_revision=2,
                file_format="md",
                file_name="phase10-d2.md",
                idempotency_key="phase10-d2-md",
            )
            return task.id, dict(task.input_ref)

    d2_task_id, d2_input = runtime.run(create_d2_task)
    assert d2_task_id not in {task_id for task_id, _ in task_formats}
    assert d2_input["body_markdown"].startswith("# D2")
    assert d2_input["tables"] == [
        {"title": "D2 表格", "columns": ["字段", "值"], "rows": [["B", "2"]]}
    ]
    assert d2_input["source_refs"] == []

    with Session(database.engine) as session:
        events = list(
            session.scalars(
                select(OutboxEvent).where(
                    OutboxEvent.aggregate_type == "generation_task",
                    OutboxEvent.aggregate_id.in_([task_id for task_id, _ in task_formats]),
                    OutboxEvent.event_type == "generation.requested",
                )
            )
        )
    event_ids = {event.aggregate_id: event.id for event in events}
    assert set(event_ids) == {task_id for task_id, _ in task_formats}
    for task_id, _ in task_formats:
        assert runtime.run(runtime.worker.generation.handle_event, event_ids[task_id])

    def extracted_content(file_format: str, payload: bytes) -> str:
        if file_format == "md":
            return payload.decode("utf-8")
        if file_format == "docx":
            from docx import Document as WordDocument

            document = WordDocument(io.BytesIO(payload))
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend(
                cell.text for table in document.tables for row in table.rows for cell in row.cells
            )
            return "\n".join(parts)
        if file_format == "pdf":
            import pymupdf

            document = pymupdf.open(stream=payload, filetype="pdf")
            try:
                return "\n".join(page.get_text() for page in document)
            finally:
                document.close()
        if file_format == "xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            try:
                values: list[str] = []
                for sheet in workbook.worksheets:
                    values.append(sheet.title)
                    values.extend(
                        str(value)
                        for row in sheet.iter_rows(values_only=True)
                        for value in row
                        if value is not None
                    )
                return "\n".join(values)
            finally:
                workbook.close()
        raise AssertionError(file_format)

    with Session(database.engine) as session:
        task_rows = list(
            session.scalars(
                select(GenerationTask).where(
                    GenerationTask.id.in_([task_id for task_id, _ in task_formats])
                )
            )
        )
        generated = list(
            session.scalars(select(GeneratedFile).where(GeneratedFile.task_id.in_(event_ids)))
        )
        assert {file.file_format for file in generated} == {
            "md",
            "docx",
            "pdf",
            "xlsx",
        }, [(task.id, task.status, task.error_code, task.error_message) for task in task_rows]
        for file in generated:
            task = session.get(GenerationTask, file.task_id)
            source = session.get(StorageObject, file.storage_object_id)
            assert task is not None and source is not None
            assert task.status == "succeeded"
            assert task.input_ref["body_markdown"].startswith("# D1")
            assert task.input_ref["tables"] == [
                {"title": "D1 表格", "columns": ["字段", "值"], "rows": [["A", "1"]]}
            ]
            assert file.source_refs == frozen["source_refs"]
            assert source.content_sha256 == file.content_sha256
            destination = tmp_path / file.file_name
            runtime.worker.storage.download(
                source.object_key, destination, source.size_bytes, source.content_sha256 or ""
            )
            text = extracted_content(file.file_format, destination.read_bytes()).replace(
                "\xa0", " "
            )
            assert "D1 正文" in text and "D1 表格" in text
            assert "D2 正文" not in text and "D2 表格" not in text


def test_generated_import_requires_ready_task_and_replays_confirmation_once(runtime, database):
    content = b"# Phase 10 generated content\n\nFrozen body.\n"
    digest = hashlib.sha256(content).hexdigest()

    async def create_generated(
        status: str = "succeeded", storage_status: str = "available"
    ) -> tuple[int, int]:
        async with runtime.worker.sessions() as session, session.begin():
            owner = await session.scalar(select(User).where(User.username == "初始管理员"))
            assert owner is not None
            key = runtime.worker.storage.new_key("generated")
            await asyncio.to_thread(
                runtime.worker.storage.upload, key, io.BytesIO(content), "text/markdown"
            )
            obj = StorageObject(
                bucket_name=runtime.worker.storage.bucket,
                object_key=key,
                object_type="generated",
                size_bytes=len(content),
                content_type="text/markdown",
                content_sha256=digest,
                status=storage_status,
            )
            session.add(obj)
            await session.flush()
            task = GenerationTask(
                owner_user_id=owner.id,
                kind="content_render",
                format="md",
                idempotency_key=uuid.uuid4().hex,
                request_fingerprint="b" * 64,
                input_ref={"draft_revision": 1},
                source_refs=[],
                status=status,
                file_id=obj.id,
            )
            session.add(task)
            await session.flush()
            generated = GeneratedFile(
                task_id=task.id,
                owner_user_id=owner.id,
                storage_object_id=obj.id,
                file_name="phase10-generated.md",
                file_format="md",
                generation_kind="content_render",
                source_refs=[],
                content_sha256=digest,
                size_bytes=len(content),
            )
            session.add(generated)
            await session.flush()
            return generated.id, task.id

    not_ready, _ = runtime.run(create_generated, "queued")
    rejected = runtime.client.post(
        f"/api/v1/generated-files/{not_ready}/import-previews",
        headers={**runtime.admin, "Idempotency-Key": "phase10-not-ready"},
        json={"knowledge_base_id": runtime.base_id, "name": "not-ready.md"},
    )
    assert rejected.status_code == 404, rejected.text

    unavailable, _ = runtime.run(create_generated, "succeeded", "deleted")
    unavailable_response = runtime.client.post(
        f"/api/v1/generated-files/{unavailable}/import-previews",
        headers={**runtime.admin, "Idempotency-Key": "phase10-object-unavailable"},
        json={"knowledge_base_id": runtime.base_id, "name": "unavailable.md"},
    )
    assert unavailable_response.status_code == 503, unavailable_response.text

    file_id, task_id = runtime.run(create_generated)
    _create_user(runtime.client, runtime.admin, "phase10-generated-other")
    other_token = auth(runtime.client, "phase10-generated-other", "Phase10-Memory-Contract-2026!")
    spoofed = runtime.client.post(
        f"/api/v1/generated-files/{file_id}/import-previews",
        headers={**other_token, "Idempotency-Key": "phase10-import-spoofed-owner"},
        json={"knowledge_base_id": runtime.base_id, "name": "spoofed.md"},
    )
    assert spoofed.status_code == 404, spoofed.text

    revoked_preview = runtime.client.post(
        f"/api/v1/generated-files/{file_id}/import-previews",
        headers={**runtime.admin, "Idempotency-Key": "phase10-import-revoked-preview"},
        json={"knowledge_base_id": runtime.base_id, "name": "revoked-target.md"},
    )
    assert revoked_preview.status_code == 202, revoked_preview.text
    revoked_action = revoked_preview.json()["data"]

    async def recycle_target() -> None:
        from enterprise_mind.knowledge import lifecycle

        async with runtime.worker.sessions() as session, session.begin():
            owner = await session.scalar(select(User).where(User.username == "初始管理员"))
            assert owner is not None
            identity = CurrentIdentity(owner, None, owner.credential_version)
            target = await lifecycle.resource(
                session, identity, "knowledge_base", int(runtime.base_id), recycled=False
            )
            await lifecycle.soft_delete(session, target)

    async def restore_target() -> None:
        from enterprise_mind.knowledge import lifecycle

        async with runtime.worker.sessions() as session, session.begin():
            owner = await session.scalar(select(User).where(User.username == "初始管理员"))
            assert owner is not None
            identity = CurrentIdentity(owner, None, owner.credential_version)
            target = await lifecycle.resource(
                session, identity, "knowledge_base", int(runtime.base_id), recycled=True
            )
            await lifecycle.restore(
                session,
                identity,
                "knowledge_base",
                int(runtime.base_id),
                {"row_version": target.row_version},
            )

    runtime.run(recycle_target)
    revoked_response = runtime.client.post(
        f"/api/v1/pending-actions/{revoked_action['id']}/responses",
        headers={**runtime.admin, "Idempotency-Key": "phase10-import-revoked-response"},
        json={"decision": "accept"},
    )
    assert revoked_response.status_code == 404, revoked_response.text
    assert revoked_response.json()["error"]["code"] == ErrorCode.RESOURCE_NOT_FOUND.value
    runtime.run(restore_target)

    stale_preview = runtime.client.post(
        f"/api/v1/generated-files/{file_id}/import-previews",
        headers={**runtime.admin, "Idempotency-Key": "phase10-import-stale-preview"},
        json={"knowledge_base_id": runtime.base_id, "name": "stale-target.md"},
    )
    assert stale_preview.status_code == 202, stale_preview.text
    stale_action = stale_preview.json()["data"]
    current_base = runtime.client.get(
        f"/api/v1/knowledge-bases/{runtime.base_id}", headers=runtime.admin
    )
    assert current_base.status_code == 200, current_base.text
    renamed = runtime.client.patch(
        f"/api/v1/knowledge-bases/{runtime.base_id}",
        headers=runtime.admin,
        json={
            "row_version": current_base.json()["data"]["row_version"],
            "name": "测试知识库更新",
        },
    )
    assert renamed.status_code == 200, renamed.text
    stale_response = runtime.client.post(
        f"/api/v1/pending-actions/{stale_action['id']}/responses",
        headers={**runtime.admin, "Idempotency-Key": "phase10-import-stale-response"},
        json={"decision": "accept"},
    )
    assert stale_response.status_code == 409, stale_response.text
    assert stale_response.json()["error"]["code"] == "PENDING_ACTION_BINDING_CHANGED"

    preview = runtime.client.post(
        f"/api/v1/generated-files/{file_id}/import-previews",
        headers={**runtime.admin, "Idempotency-Key": "phase10-import-preview"},
        json={"knowledge_base_id": runtime.base_id, "name": "imported-generated.md"},
    )
    assert preview.status_code == 202, preview.text
    action = preview.json()["data"]
    assert action["preview"]["generated_file_id"] == str(file_id)
    assert action["preview"]["source_sha256"] == digest
    assert action["preview"]["source_size_bytes"] == len(content)
    assert action["preview"]["target"]["knowledge_base_id"] == str(runtime.base_id)
    replay = runtime.client.post(
        f"/api/v1/generated-files/{file_id}/import-previews",
        headers={**runtime.admin, "Idempotency-Key": "phase10-import-preview"},
        json={"knowledge_base_id": runtime.base_id, "name": "imported-generated.md"},
    )
    assert replay.status_code == 202
    assert replay.json()["data"]["id"] == action["id"]
    with Session(database.engine) as session:
        assert (
            session.scalar(select(Document.id).where(Document.name == "imported-generated.md"))
            is None
        )

    response = runtime.client.post(
        f"/api/v1/pending-actions/{action['id']}/responses",
        headers={**runtime.admin, "Idempotency-Key": "phase10-import-response"},
        json={"decision": "accept"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["data"]["status"] == "accepted"
    replay_response = runtime.client.post(
        f"/api/v1/pending-actions/{action['id']}/responses",
        headers={**runtime.admin, "Idempotency-Key": "phase10-import-response"},
        json={"decision": "accept"},
    )
    assert replay_response.status_code == 200

    # Run the real parse event. The isolated core fixture has no embedding/index
    # provider, so overwrite must still be refused while its processing task is
    # pending instead of pretending that indexing succeeded.
    assert runtime.run(runtime.dispatcher.dispatch_once) == 1
    assert runtime.run(runtime.worker.run_once)

    overwrite_preview = runtime.client.post(
        f"/api/v1/generated-files/{file_id}/import-previews",
        headers={**runtime.admin, "Idempotency-Key": "phase10-import-overwrite-preview"},
        json={
            "knowledge_base_id": runtime.base_id,
            "name": "imported-generated.md",
            "conflict_action": "overwrite",
        },
    )
    assert overwrite_preview.status_code == 202, overwrite_preview.text
    overwrite_action = overwrite_preview.json()["data"]
    assert overwrite_action["operation_name"] == "document.overwrite"
    overwrite_response = runtime.client.post(
        f"/api/v1/pending-actions/{overwrite_action['id']}/responses",
        headers={
            **runtime.admin,
            "Idempotency-Key": "phase10-import-overwrite-response",
        },
        json={"decision": "accept"},
    )
    assert overwrite_response.status_code == 409, overwrite_response.text
    assert (
        overwrite_response.json()["error"]["code"]
        == ErrorCode.DOCUMENT_PROCESSING_IN_PROGRESS.value
    )

    same_name = runtime.client.post(
        f"/api/v1/generated-files/{file_id}/import-previews",
        headers={**runtime.admin, "Idempotency-Key": "phase10-import-name-conflict"},
        json={
            "knowledge_base_id": runtime.base_id,
            "name": "imported-generated.md",
            "conflict_action": "reject",
        },
    )
    assert same_name.status_code == 409, same_name.text
    assert same_name.json()["error"]["code"] == ErrorCode.DOCUMENT_NAME_CONFLICT.value

    expired_file, _ = runtime.run(create_generated)
    expired_preview = runtime.client.post(
        f"/api/v1/generated-files/{expired_file}/import-previews",
        headers={**runtime.admin, "Idempotency-Key": "phase10-import-expired-preview"},
        json={"knowledge_base_id": runtime.base_id, "name": "expired-generated.md"},
    )
    assert expired_preview.status_code == 202, expired_preview.text
    expired_action = expired_preview.json()["data"]
    with Session(database.engine) as session, session.begin():
        pending = session.get(PendingAction, int(expired_action["id"]))
        assert pending is not None
        pending.created_at = datetime.now(UTC) - timedelta(hours=1)
        pending.expires_at = datetime.now(UTC) - timedelta(seconds=1)
    expired_response = runtime.client.post(
        f"/api/v1/pending-actions/{expired_action['id']}/responses",
        headers={**runtime.admin, "Idempotency-Key": "phase10-import-expired-response"},
        json={"decision": "accept"},
    )
    assert expired_response.status_code == 409
    assert expired_response.json()["error"]["code"] == "PENDING_ACTION_EXPIRED"

    with Session(database.engine) as session:
        documents = list(
            session.scalars(select(Document).where(Document.name == "imported-generated.md"))
        )
        assert len(documents) == 1
        assert documents[0].knowledge_base_id == int(runtime.base_id)
        # The real parser worker has claimed this revision by this point.  The
        # core fixture intentionally has no embedding/index provider, so the
        # document may still be pending on an unclaimed run or be processing
        # after the parse task is claimed; both states prove no overwrite
        # revision was created by the rejected conflict response.
        assert documents[0].processing_status in {"pending", "processing"}
        revision = session.scalar(
            select(Revision)
            .where(Revision.document_id == documents[0].id)
            .order_by(Revision.revision_no.desc())
        )
        assert revision is not None
        assert revision.revision_no == 1 and revision.trigger_type == "generated_write"
        assert revision.processing_status in {"pending", "processing"}
        assert revision.content_sha256 == digest and revision.size_bytes == len(content)
        copied = session.get(StorageObject, revision.original_object_id)
        assert copied is not None
        assert copied.content_sha256 == digest and copied.size_bytes == len(content)
        events = list(
            session.scalars(
                select(OutboxEvent).where(OutboxEvent.event_type == "document.processing.requested")
            )
        )
        assert any(event.payload.get("document_id") == str(documents[0].id) for event in events)
        task = session.get(GenerationTask, task_id)
        assert task is not None and task.status == "succeeded"
