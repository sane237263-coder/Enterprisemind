"""Run inside the isolated converter image; use generated non-sensitive Office samples."""

import argparse
import asyncio
import json
import os
import secrets
import subprocess
import sys
import time
from pathlib import Path
from tempfile import TemporaryDirectory

import httpx
from docx import Document
from enterprise_mind.infrastructure.settings import OfficeConverterConfig, WorkerConfig
from enterprise_mind.ingestion.office import OfficeConverter
from enterprise_mind.ingestion.parsing import parse_file
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    version = subprocess.check_output(["libreoffice", "--version"], text=True).strip()
    if not version.startswith("LibreOffice 26.2.6"):
        raise RuntimeError("The deployment baseline requires LibreOffice 26.2.6")
    with TemporaryDirectory(prefix="em-converter-check-") as directory:
        root = Path(directory)
        secret = root / "token"
        secret.write_text(secrets.token_urlsafe(32))
        os.environ["EM_OFFICE_CONVERTER_TOKEN_FILE"] = str(secret)
        document = Document()
        document.add_heading("ABC-123 企业规范", 1)
        document.add_paragraph("支持企业加密。")
        document.save(root / "sample.docx")
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(5), Inches(1)
        ).text = "ABC-123 设备架构"
        presentation.save(root / "sample.pptx")
        workbook = Workbook()
        workbook.active.append(["型号", "数量"])
        workbook.active.append(["ABC-123", 3])
        workbook.save(root / "sample.xlsx")
        legacy = root / "legacy"
        legacy.mkdir()
        for modern, target in [
            ("docx", "doc:MS Word 97"),
            ("pptx", "ppt:MS PowerPoint 97"),
            ("xlsx", "xls:MS Excel 97"),
        ]:
            subprocess.run(
                [
                    "libreoffice",
                    "-env:UserInstallation=" + (root / "seed-profile").as_uri(),
                    "--headless",
                    "--convert-to",
                    target,
                    "--outdir",
                    str(legacy),
                    str(root / ("sample." + modern)),
                ],
                check=True,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                timeout=60,
            )
        service = subprocess.Popen(
            [
                sys.executable,
                "-m",
                "uvicorn",
                "enterprise_mind.office_converter:factory",
                "--factory",
                "--host",
                "127.0.0.1",
                "--port",
                "18808",
                "--no-access-log",
            ],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        try:
            for _ in range(100):
                try:
                    if (
                        httpx.get(
                            "http://127.0.0.1:18808/health/ready",
                            trust_env=False,
                            timeout=8,
                        ).status_code
                        == 200
                    ):
                        break
                except httpx.TransportError:
                    pass
                if service.poll() is not None:
                    raise RuntimeError("Converter startup failed")
                time.sleep(0.1)
            else:
                raise RuntimeError("Converter readiness timed out")
            converter = OfficeConverter(
                OfficeConverterConfig(base_url="http://127.0.0.1:18808"), 104857600
            )
            results = []
            for extension, normalized in [
                ("doc", "docx"),
                ("ppt", "pptx"),
                ("xls", "xlsx"),
            ]:
                target = root / ("result." + normalized)
                actual = asyncio.run(
                    converter.convert(
                        legacy / ("sample." + extension), extension, target
                    )
                )
                assert actual == normalized
                parsed = parse_file(
                    target, normalized, WorkerConfig(), root / extension
                )
                assert "ABC-123" in "\n".join(s.text for s in parsed.sources)
                assert (root / extension / "preview.html").is_file()
                results.append(
                    {
                        "source": extension,
                        "normalized": normalized,
                        "text_and_preview": "passed",
                    }
                )
            args.output.parent.mkdir(parents=True, exist_ok=True)
            args.output.write_text(
                json.dumps(
                    {
                        "converter": version,
                        "network": "loopback only",
                        "formats": results,
                    },
                    ensure_ascii=False,
                    indent=2,
                ),
                encoding="utf-8",
            )
            print("Real LibreOffice conversion and HTTP client passed: DOC, PPT, XLS.")
        finally:
            service.terminate()
            service.wait(timeout=15)


if __name__ == "__main__":
    main()
